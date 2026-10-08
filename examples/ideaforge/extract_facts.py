"""Extract sourced facts from the ideaForge company-profile deck into facts.json.

Every value is read from the deck's own tables, charts and text by pattern; nothing is
typed by hand. Each fact cites the filing or data source the deck cites for it, plus the
deck page. Figures that appear in more than one place are cross-checked.

    uv run python examples/ideaforge/extract_facts.py SOURCE.pptx examples/ideaforge/facts.json
"""

from __future__ import annotations

import re
import statistics
import sys
from decimal import Decimal
from pathlib import Path
from typing import Any

from pptx import Presentation

from deckforge.ingest.models import Fact, FactSet, Unit

MILLION = Decimal(1_000_000)
BILLION = Decimal(1_000_000_000)
DECK = "ideaForge company profile deck"
IDEAFORGE = ("ideaforge", "ideaForge Technology Ltd.")

# Per-column sources for the P&L table, as the deck's slide 4 footnote assigns them.
PNL_SOURCES = {
    "fy23": "ideaForge Q4 FY24 investor release (BSE, 14-May-24)",
    "fy24": "ideaForge Q4 FY24 investor release (BSE, 14-May-24)",
    "fy25": "ideaForge Q4 FY25 results release (May-25)",
    "fy26": "ideaForge Q3 FY26 investor presentation (NSE, 22-Jan-26) for 9M FY26 plus "
    "ideaForge Q1 FY27 investor presentation (NSE, 10-Aug-26) for Q4 FY26",
    "q1fy26": "ideaForge Q1 FY27 investor presentation (NSE, 10-Aug-26)",
    "q1fy27": "ideaForge Q1 FY27 investor presentation (NSE, 10-Aug-26)",
    "ltm_jun26": "ideaForge Q1 FY27 investor presentation (NSE, 10-Aug-26)",
}
PNL_ROWS = {
    "Revenue from operations": ("revenue", Unit.INR),
    "Growth (YoY)": ("revenue_growth", Unit.PERCENT),
    "Gross profit": ("gross_profit", Unit.INR),
    "Gross margin": ("gross_margin", Unit.PERCENT),
    "o/w other income (in EBITDA)": ("other_income", Unit.INR),
    "EBITDA (company-reported)": ("ebitda", Unit.INR),
    "EBITDA margin": ("ebitda_margin", Unit.PERCENT),
    "PAT": ("pat", Unit.INR),
    "PAT margin": ("pat_margin", Unit.PERCENT),
}
COMPANIES = {
    "Paras Defence (IN)": ("parasdefence", "Paras Defence and Space Technologies"),
    "Data Patterns (IN)": ("datapatterns", "Data Patterns (India)"),
    "Zen Technologies (IN)": ("zentech", "Zen Technologies"),
    "AeroVironment (US)": ("aerovironment", "AeroVironment"),
    "ideaForge (IN)": IDEAFORGE,
}
COMPS_COLUMNS = {
    "Mkt cap (m)": "market_cap",
    "EV (m)": "ev",
    "LTM revenue (m)": "revenue_ltm",
    "EBITDA margin": "ebitda_margin_ltm",
    "EV / revenue": "ev_revenue_ltm",
    "EV / EBITDA": "ev_ebitda_ltm",
    "P / E": "pe_ltm",
}
VALUATION_BARS = {
    "52-week trading range": "trading_52w",
    "EV / LTM revenue": "ev_revenue",
    "EV / LTM EBITDA (ex-other income)": "ev_ebitda",
}
NUMBER = re.compile(r"^\(?([₹$])?\(?(-?[\d,]+(?:\.\d+)?)\)?(%|x)?\)?$")


class ExtractionError(ValueError):
    """The deck does not contain a value where the extractor expects one."""


def parse(text: str) -> tuple[Decimal, str | None, str | None] | None:
    """'₹(315.3)' -> (-315.3, '₹', None); '(48.7%)' -> (-48.7, None, '%'); '-'/'n.a.' -> None."""
    t = text.strip().replace("\u2212", "-")
    m = NUMBER.match(t)
    if not m:
        return None
    value = Decimal(m[2].replace(",", ""))
    negative = t.startswith("(") or "(" in t[:2]
    return (-value if negative else value), m[1], m[3]


class Deck:
    def __init__(self, path: Path) -> None:
        self.prs = Presentation(str(path))
        self.facts: dict[str, Fact] = {}

    def _shape(self, n: int, name: str) -> Any:
        for shape in self.prs.slides[n - 1].shapes:
            if shape.name == name:
                return shape
        raise ExtractionError(f"slide {n} has no shape named {name!r}")

    def text(self, n: int) -> str:
        slide = self.prs.slides[n - 1]
        return "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)

    def sources(self, n: int) -> str:
        """The slide's own 'Sources:' footnote, as cited by the deck."""
        return self._shape(n, "Source").text_frame.text.removeprefix("Sources: ").strip()

    def table(self, n: int, name: str) -> list[list[str]]:
        tbl = self._shape(n, name).table
        return [[c.text.strip() for c in row.cells] for row in tbl.rows]

    def chart(self, n: int, name: str) -> tuple[list[str], dict[str, list[Decimal]]]:
        plot = self._shape(n, name).chart.plots[0]
        series = {s.name: [Decimal(str(v)) for v in s.values] for s in plot.series}
        return [str(c) for c in plot.categories], series

    def grab(self, n: int, pattern: str) -> str:
        found = re.findall(pattern, self.text(n))
        if len(set(found)) != 1:
            raise ExtractionError(f"slide {n}: expected one match for {pattern!r}, got {found}")
        return str(found[0])

    def add(
        self,
        entity: tuple[str, str],
        metric: str,
        period: str,
        value: Decimal,
        unit: Unit,
        source: str,
        page: int,
        label: str | None = None,
    ) -> str:
        fact_id = f"{entity[0]}.{metric}.{period.lower().replace(' ', '_').replace('-', '_')}"
        fact = Fact(
            id=fact_id,
            entity=entity[1],
            metric=metric,
            period=label or period,
            value=value,
            unit=unit,
            source_ref=f"{source} [{DECK}, p.{page}]",
        )
        existing = self.facts.get(fact_id)
        if existing is not None and existing.value != value:
            raise ExtractionError(f"{fact_id}: {existing.value} vs {value}")
        self.facts[fact_id] = fact
        return fact_id

    def value(self, fact_id: str) -> Decimal:
        return self.facts[fact_id].value


def money(text: str, scale: Decimal = MILLION) -> Decimal:
    parsed = parse(text)
    if parsed is None:
        raise ExtractionError(f"not a number: {text!r}")
    return parsed[0] * scale


def pct(text: str) -> Decimal:
    return money(text, Decimal("0.01"))


def extract(deck: Deck) -> FactSet:
    # --- p.4 P&L table: per-column sources --------------------------------------------
    rows = deck.table(4, "P&L table")
    periods = ["fy23", "fy24", "fy25", "fy26", "q1fy26", "q1fy27", "ltm_jun26"]
    labels = {
        "fy23": "FY23",
        "fy24": "FY24",
        "fy25": "FY25",
        "fy26": "FY26",
        "q1fy26": "Q1 FY26",
        "q1fy27": "Q1 FY27",
        "ltm_jun26": "LTM Jun-26",
    }
    if rows[0][1:] != [labels[p] for p in periods]:
        raise ExtractionError(f"unexpected P&L header {rows[0]}")
    for row in rows[1:]:
        metric, unit = PNL_ROWS[row[0]]
        for period, cell in zip(periods, row[1:], strict=True):
            parsed = parse(cell)
            if parsed is None:  # a dash or 'n.a.': not disclosed, no fact
                continue
            value = parsed[0] * (MILLION if unit is Unit.INR else Decimal("0.01"))
            deck.add(IDEAFORGE, metric, period, value, unit, PNL_SOURCES[period], 4)
    q1_src = PNL_SOURCES["q1fy27"]
    cats, series = deck.chart(4, "Quarterly chart")
    for cat, rev, ebitda in zip(cats, series["Revenue"], series["EBITDA"], strict=True):
        period = cat.lower().replace(" ", "")
        deck.add(IDEAFORGE, "revenue", period, rev * MILLION, Unit.INR, q1_src, 4)
        deck.add(IDEAFORGE, "ebitda", period, ebitda * MILLION, Unit.INR, q1_src, 4)
    cats, series = deck.chart(4, "Order book chart")
    for cat, ob in zip(cats, series["Order book"], strict=True):
        deck.add(
            IDEAFORGE, "order_book", cat.lower().replace("-", ""), ob * MILLION, Unit.INR, q1_src, 4
        )
    deck.add(
        IDEAFORGE,
        "revenue",
        "q1fy25",
        money(deck.grab(4, r"Q1 FY25 revenue of (₹[\d,.]+)m")),
        Unit.INR,
        "ideaForge Q1 FY26 results release",
        4,
    )

    # --- p.2 takeaways: slide-level sources ------------------------------------------
    src2 = deck.sources(2).split(". Price")[0]
    deck.add(
        IDEAFORGE,
        "order_book_executed_share",
        "q1fy27",
        pct(deck.grab(2, r"(\d+)%\+ of the opening order book executed")),
        Unit.PERCENT,
        src2,
        2,
    )
    deck.add(
        IDEAFORGE,
        "qip_proceeds",
        "jul26",
        money(deck.grab(2, r"(₹[\d,.]+)m QIP at")),
        Unit.INR,
        src2,
        2,
    )
    deck.add(
        IDEAFORGE,
        "qip_price",
        "jul26",
        money(deck.grab(2, r"QIP at (₹[\d,.]+)"), Decimal(1)),
        Unit.INR_PER_SHARE,
        src2,
        2,
    )
    deck.add(
        IDEAFORGE,
        "rdi_funding",
        "fy27",
        money(deck.grab(2, r"(₹[\d,.]+)m of low-cost funding")),
        Unit.INR,
        src2,
        2,
    )
    deck.add(
        ("india_mod", "Ministry of Defence (India)"),
        "drone_procurement_fast_track",
        "plan",
        money(deck.grab(2, r"plans (₹[\d,.]+)bn of drone"), BILLION),
        Unit.INR,
        src2,
        2,
    )
    deck.add(
        ("india_dac", "Defence Acquisition Council (India)"),
        "capital_acquisitions",
        "approved",
        money(deck.grab(2, r"DAC approved (₹[\d,.]+)bn"), BILLION),
        Unit.INR,
        src2,
        2,
    )
    price_src = deck.grab(2, r"Price ₹[\d.]+ \((NSE close, [^)]+)\)")
    deck.add(
        IDEAFORGE,
        "share_price",
        "06-Oct-26",
        money(deck.grab(2, r"Price (₹[\d.]+) \(NSE"), Decimal(1)),
        Unit.INR_PER_SHARE,
        price_src,
        2,
    )

    # --- p.3 company facts ------------------------------------------------------------
    src3 = deck.sources(3)
    facts_tbl = {r[0]: r[1] for r in deck.table(3, "Key facts table")}
    rank = re.match(r"#(\d+) in dual-use drones \(DRONEII, ([^)]+)\)", facts_tbl["Global rank"])
    missions = re.match(r"([\d,]+)\+ customer missions", facts_tbl["Field record"])
    ip = re.match(r"(\d+) patents: (\d+) granted, (\d+) pending", facts_tbl["IP"])
    shares = re.match(r"([\d.]+)m post-QIP", facts_tbl["Shares out"])
    if not (rank and missions and ip and shares):
        raise ExtractionError(f"key facts table changed: {facts_tbl}")
    deck.add(
        IDEAFORGE,
        "dual_use_drone_rank",
        "dec24",
        Decimal(rank[1]),
        Unit.COUNT,
        f"DRONEII ranking ({rank[2]})",
        3,
    )
    deck.add(
        IDEAFORGE, "customer_missions", "jul26", money(missions[1], Decimal(1)), Unit.COUNT, src3, 3
    )
    for key, grp in (("patents_total", 1), ("patents_granted", 2), ("patents_pending", 3)):
        deck.add(IDEAFORGE, key, "jun26", Decimal(ip[grp]), Unit.COUNT, src3, 3)
    deck.add(
        IDEAFORGE,
        "shares_outstanding",
        "post_qip",
        Decimal(shares[1]) * MILLION,
        Unit.SHARES,
        src3,
        3,
    )
    cats, series = deck.chart(3, "Revenue mix chart")
    for cat, share in zip(cats, next(iter(series.values())), strict=True):
        deck.add(
            IDEAFORGE, f"revenue_mix_{cat.lower()}", "q1fy27", share / 100, Unit.PERCENT, src3, 3
        )
    deck.add(
        IDEAFORGE,
        "tech_workforce_share",
        "fy27",
        pct(deck.grab(3, r"(\d+)%\+ of workforce")),
        Unit.PERCENT,
        src3,
        3,
    )
    payload = re.search(
        r"(\d+)\u2013(\d+) kg", "\n".join(" ".join(r) for r in deck.table(3, "Portfolio table"))
    )
    if not payload:
        raise ExtractionError("YETI payload range not found")
    deck.add(
        ("ideaforge_yeti", "ideaForge YETI"),
        "payload_min_kg",
        "trl6",
        Decimal(payload[1]),
        Unit.COUNT,
        src3,
        3,
    )
    deck.add(
        ("ideaforge_yeti", "ideaForge YETI"),
        "payload_max_kg",
        "trl6",
        Decimal(payload[2]),
        Unit.COUNT,
        src3,
        3,
    )

    # --- p.5 trading comps --------------------------------------------------------------
    comps = deck.table(5, "Comps table")
    header = comps[0]
    peer_src = "stockanalysis.com statistics page (S&P Global Market Intelligence data)"
    if_src = (
        "ideaForge Q1 FY27 investor presentation (NSE, 10-Aug-26); ideaForge audited "
        "consolidated results FY26 (NSE, 30-Apr-26); QIP allotment filing (Jul-26)"
    )
    median_row: list[str] | None = None
    for row in comps[1:]:
        if row[0] == "Indian peer median":
            median_row = row
            continue
        entity = COMPANIES[row[0]]
        price_date = row[header.index("Price date")]
        source = if_src if entity is IDEAFORGE else peer_src
        for col, metric in COMPS_COLUMNS.items():
            parsed = parse(row[header.index(col)])
            if parsed is None:  # NM
                continue
            number, symbol, suffix = parsed
            if suffix == "%":
                value, unit = number / 100, Unit.PERCENT
            elif suffix == "x":
                value, unit = number, Unit.RATIO
            else:
                value, unit = number * MILLION, Unit.USD if symbol == "$" else Unit.INR
            if entity is IDEAFORGE and metric == "revenue_ltm":
                if value != deck.value("ideaforge.revenue.ltm_jun26"):
                    raise ExtractionError("comps LTM revenue disagrees with the P&L table")
                continue  # one fact per number: reuse the P&L LTM revenue
            deck.add(
                entity, metric, "ltm", value, unit, source, 5, label=f"LTM; price {price_date}"
            )
    if median_row is None:
        raise ExtractionError("comps table has no Indian peer median row")
    indian = ["parasdefence", "datapatterns", "zentech"]
    for col in ("EV / revenue", "EV / EBITDA"):
        metric = COMPS_COLUMNS[col]
        stated = money(median_row[header.index(col)], Decimal(1))
        computed = statistics.median(deck.value(f"{p}.{metric}.ltm") for p in indian)
        if stated != computed:
            raise ExtractionError(f"{col} median {stated} != median of peers {computed}")
    deck.add(
        IDEAFORGE,
        "ev_revenue_discount_to_indian_median",
        "ltm",
        pct(deck.grab(5, r"is a ([\d.]+)% discount")),
        Unit.PERCENT,
        "deck calculation from the trading comps",
        5,
    )
    deck.add(
        IDEAFORGE,
        "net_cash_pro_forma",
        "31-mar-26_plus_qip",
        money(deck.grab(5, r"pro forma net cash (₹[\d,.]+)m")),
        Unit.INR,
        if_src,
        5,
    )

    # --- p.6 football field -----------------------------------------------------------
    note6 = deck._shape(6, "Source").text_frame.text
    ex_oi = money(re.search(r"LTM EBITDA ex-other income (₹[\d,.]+)m", note6)[1])  # type: ignore[index]
    if ex_oi != deck.value("ideaforge.ebitda.ltm_jun26") - deck.value(
        "ideaforge.other_income.ltm_jun26"
    ):
        raise ExtractionError("EBITDA ex-other income does not tie to the P&L table")
    deck.add(
        IDEAFORGE,
        "ebitda_ex_other_income",
        "ltm_jun26",
        ex_oi,
        Unit.INR,
        PNL_SOURCES["ltm_jun26"],
        6,
    )
    cats, series = deck.chart(6, "Football field chart")
    valuation_src = (
        "deck calculation: (peer multiple x LTM metric + pro forma net cash) / "
        "post-QIP shares; peer multiples per stockanalysis.com"
    )
    week52_src = deck.grab(6, r"52-week range per (Google Finance \([^)]+\))")
    text6 = deck.text(6)
    for cat, low, span in zip(cats, series["Low"], series["Range"], strict=True):
        key = VALUATION_BARS[cat]
        src = week52_src if key == "trading_52w" else valuation_src
        high = low + span
        for side, v in (("low", low), ("high", high)):
            shown = f"₹{v:,.1f}"
            if shown not in text6:
                raise ExtractionError(f"{cat} {side} {shown} not shown on slide 6")
            deck.add(
                IDEAFORGE, f"implied_value_{key}_{side}", "06-Oct-26", v, Unit.INR_PER_SHARE, src, 6
            )
    return FactSet(facts=sorted(deck.facts.values(), key=lambda f: f.id))


def main(argv: list[str]) -> int:
    source, out = Path(argv[1]), Path(argv[2])
    facts = extract(Deck(source))
    out.parent.mkdir(parents=True, exist_ok=True)
    facts.dump(out)
    print(f"{len(facts.facts)} facts -> {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
