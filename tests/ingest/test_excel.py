from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from deckforge.ingest import excel
from deckforge.ingest.models import FactSet, Unit


def test_template_roundtrip_reports_gaps(tmp_path: Path) -> None:
    path = excel.write_template(
        tmp_path / "inputs.xlsx",
        [
            {
                "id": "testco.share_price.current",
                "entity": "Testco",
                "metric": "share_price",
                "period": "current",
                "unit": "USD_per_share",
            },
            {
                "id": "testco.val_dcf.low",
                "entity": "Testco",
                "metric": "val_dcf_low",
                "period": "current",
                "unit": "USD_per_share",
            },
        ],
    )
    wb = load_workbook(path)
    wb["facts"]["E2"] = 42.5
    wb.save(path)
    result = excel.read_facts(path)
    assert result.missing == ["testco.val_dcf.low"]
    (fact,) = result.facts
    assert fact.value == Decimal("42.5") and fact.unit is Unit.USD_PER_SHARE
    assert fact.source_ref == "inputs.xlsx!facts!E2"


def test_bad_header(tmp_path: Path) -> None:
    path = excel.write_template(tmp_path / "x.xlsx", [])
    wb = load_workbook(path)
    wb["facts"]["A1"] = "name"
    wb.save(path)
    with pytest.raises(excel.ExcelFactsError, match="header"):
        excel.read_facts(path)


def test_merge_conflicts() -> None:
    base = {"entity": "T", "metric": "m", "period": "p", "unit": "USD", "source_ref": "s"}
    a = FactSet.model_validate({"facts": [{**base, "id": "a", "value": "1"}]})
    b = FactSet.model_validate(
        {"facts": [{**base, "id": "a", "value": "1"}, {**base, "id": "b", "value": "2"}]}
    )
    assert [f.id for f in a.merge(b).facts] == ["a", "b"]
    c = FactSet.model_validate({"facts": [{**base, "id": "a", "value": "3"}]})
    with pytest.raises(ValueError, match="conflicting"):
        a.merge(c)
