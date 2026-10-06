from __future__ import annotations

from pptx.presentation import Presentation

from deckforge.qc.checkers import sources
from tests.qc.checkers.helpers import add_box
from tests.qc.conftest import MakeCtx


def test_clean_deck_passes(make_ctx: MakeCtx) -> None:
    assert sources.check(make_ctx()) == []


def test_typed_number_is_flagged(make_ctx: MakeCtx) -> None:
    ctx = make_ctx(lambda prs: add_box(prs, 2, "Margins near 45% next year"))
    assert any("not bound to facts" in i.message for i in sources.check(ctx))


def test_years_and_periods_are_not_data(make_ctx: MakeCtx) -> None:
    ctx = make_ctx(lambda prs: add_box(prs, 2, "Outlook for FY2026 and CY 2027E"))
    assert sources.check(ctx) == []


def test_missing_source_footnote(make_ctx: MakeCtx) -> None:
    def drop_source(prs: Presentation) -> None:
        for shape in list(prs.slides[3].shapes):
            if shape.has_text_frame and shape.text_frame.text.startswith("Source:"):
                shape.element.getparent().remove(shape.element)

    assert any("no 'Source:'" in i.message for i in sources.check(make_ctx(drop_source)))
