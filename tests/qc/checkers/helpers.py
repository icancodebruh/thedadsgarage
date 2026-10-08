from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.presentation import Presentation
from pptx.util import Inches, Pt


def add_box(
    prs: Presentation,
    slide: int,
    text: str,
    *,
    font: str | None = None,
    size: float | None = None,
    color: str | None = None,
    x: float = 1.0,
    w: float = 3.0,
    h: float = 0.5,
    wrap: bool = False,
) -> None:
    box = prs.slides[slide - 1].shapes.add_textbox(Inches(x), Inches(2), Inches(w), Inches(h))
    if wrap:  # fixed-size, wrapping box (python-pptx defaults to unwrapped, auto-size)
        box.text_frame.word_wrap = True
        box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    if font:
        run.font.name = font
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
