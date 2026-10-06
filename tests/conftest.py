"""Shared fixtures: a small self-made reference deck built with python-pptx."""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from deckforge.template.analyze import analyze_template, write_outputs

NAVY = RGBColor.from_string("113D63")
GREY = RGBColor.from_string("CCD1D7")


@pytest.fixture(scope="session")
def reference_deck(tmp_path_factory: pytest.TempPathFactory) -> Path:
    tmp = tmp_path_factory.mktemp("deck")
    png = tmp / "chart.png"
    Image.new("RGB", (400, 300), "white").save(png)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    title_layout, content_layout = prs.slide_layouts[0], prs.slide_layouts[5]  # title only

    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = "Project Falcon"

    titles = ["Executive Summary", "Valuation Overview for Falcon", "Trading Comparables"]
    for i, title in enumerate(titles):
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = title
        box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.5), Inches(4), Inches(1)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = NAVY
        run = box.text_frame.paragraphs[0].add_run()
        run.text = "Key message"
        run.font.name, run.font.size = "Arial", Pt(12)
        run.font.color.rgb = RGBColor.from_string("FFFFFF")
        if i == 1:
            table = slide.shapes.add_table(3, 2, Inches(0.5), Inches(3), Inches(6), Inches(1.5))
            table.table.cell(0, 0).text = "Metric"
            table.table.cell(1, 0).fill.solid()
            table.table.cell(1, 0).fill.fore_color.rgb = GREY
        if i == 2:
            slide.shapes.add_picture(str(png), Inches(0.5), Inches(3), Inches(8), Inches(3.5))

    path = tmp / "reference.pptx"
    prs.save(str(path))
    return path


# Synthetic test facts (not real company data).
TEST_SOURCE = "test fixture"
TEST_VALUES = {
    "revenue": ("USD", ["1000000000", "1250000000"]),
    "net_income": ("USD", ["-12340000", "56780000"]),
    "ebitda_margin": ("percent", ["0.2", "0.255"]),
}


def write_json(path: Path, data: object) -> Path:
    path.write_text(json.dumps(data, indent=2), encoding="utf-8")
    return path


@pytest.fixture
def deck_inputs(reference_deck: Path, tmp_path: Path) -> Path:
    """A spec directory: analyzed template, synthetic facts, and a 2-slide deck_spec.json."""
    style, library = analyze_template(reference_deck)
    write_outputs(style, library, tmp_path)
    shutil.copy(reference_deck, tmp_path / "template.pptx")
    facts = [
        {
            "id": f"testco.{metric}.fy{2024 + i}",
            "entity": "Testco",
            "metric": metric,
            "period": f"FY{2024 + i}",
            "value": v,
            "unit": unit,
            "source_ref": TEST_SOURCE,
        }
        for metric, (unit, values) in TEST_VALUES.items()
        for i, v in enumerate(values)
    ]
    write_json(tmp_path / "facts.json", {"facts": facts})
    usd = {"kind": "currency", "scale": "millions", "decimals": 1}
    spec = {
        "brief": "Test deck",
        "template": {"pptx": "template.pptx", "style": "style.json", "layouts": "layouts.json"},
        "facts": "facts.json",
        "slides": [
            {
                "slide_type": "title",
                "title": "Project Testco",
                "subtitle": "Board update",
                "date": "2026-10-06",
            },
            {
                "slide_type": "financial_table",
                "title": "Testco Financial Summary",
                "columns": ["FY 2024", "FY 2025"],
                "rows": [
                    {
                        "label": "Revenue",
                        "cells": ["testco.revenue.fy2024", "testco.revenue.fy2025"],
                        "format": usd,
                    },
                    {
                        "label": "% Margin",
                        "cells": [None, "testco.ebitda_margin.fy2025"],
                        "format": {"kind": "percent"},
                        "style": "memo",
                    },
                    {
                        "label": "Net Income",
                        "cells": ["testco.net_income.fy2024", "testco.net_income.fy2025"],
                        "format": usd,
                        "style": "total",
                    },
                ],
            },
        ],
    }
    write_json(tmp_path / "deck_spec.json", spec)
    return tmp_path
