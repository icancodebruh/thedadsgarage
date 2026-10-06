from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from tests.conftest import build_full


def test_overview_columns(full_inputs: Path) -> None:
    slide = Presentation(str(build_full(full_inputs))).slides[2]
    texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
    assert "Business Overview" in texts and "Key Statistics" in texts
    assert "Testco makes test fixtures.\nGross margin of 60.0%" in texts
    table = next(s.table for s in slide.shapes if s.has_table)
    assert [c.text for c in table.rows[0].cells] == ["Share price", "$42.50"]
    highlight = next(
        s for s in slide.shapes if s.has_text_frame and s.text_frame.text.startswith("Testco makes")
    ).text_frame.paragraphs[1]
    assert highlight._p.pPr.find("{*}buChar") is not None, "highlights use real bullets"
