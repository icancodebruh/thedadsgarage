from __future__ import annotations

import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from deckforge.build.renderer import SpecError, build_deck, numbered_layouts
from deckforge.build.slide_types.base import add_slide_number
from deckforge.ingest.models import MissingFactError
from deckforge.spec.models import DeckSpec
from deckforge.template.analyze import analyze_template


def _edit_spec(inputs: Path, **changes: object) -> DeckSpec:
    spec = json.loads((inputs / "deck_spec.json").read_text())
    spec.update(changes)
    return DeckSpec.model_validate(spec)


def test_replaces_template_slides(deck_inputs: Path, tmp_path: Path) -> None:
    out = build_deck(
        DeckSpec.load(deck_inputs / "deck_spec.json"), deck_inputs, tmp_path / "deck.pptx"
    )
    assert len(Presentation(str(out)).slides) == 2


def test_deterministic_bytes(deck_inputs: Path, tmp_path: Path) -> None:
    spec = DeckSpec.load(deck_inputs / "deck_spec.json")
    a = build_deck(spec, deck_inputs, tmp_path / "a.pptx")
    b = build_deck(spec, deck_inputs, tmp_path / "b.pptx")
    assert a.read_bytes() == b.read_bytes()


def test_missing_facts_reported_before_build(deck_inputs: Path, tmp_path: Path) -> None:
    (deck_inputs / "facts.json").write_text('{"facts": []}')
    with pytest.raises(MissingFactError) as err:
        build_deck(
            DeckSpec.load(deck_inputs / "deck_spec.json"), deck_inputs, tmp_path / "deck.pptx"
        )
    assert len(err.value.missing) == 5
    assert not (tmp_path / "deck.pptx").exists()


def test_unknown_layout(deck_inputs: Path, tmp_path: Path) -> None:
    spec = _edit_spec(
        deck_inputs, slides=[{"slide_type": "title", "title": "x", "layout_id": "nope"}]
    )
    with pytest.raises(SpecError, match="unknown layout_id"):
        build_deck(spec, deck_inputs, tmp_path / "deck.pptx")


def test_style_must_match_template(deck_inputs: Path, tmp_path: Path) -> None:
    (deck_inputs / "template.pptx").write_bytes(b"different")
    with pytest.raises(SpecError, match="re-run analyze"):
        build_deck(
            DeckSpec.load(deck_inputs / "deck_spec.json"), deck_inputs, tmp_path / "deck.pptx"
        )


def test_numbered_layouts_learned_from_reference(reference_deck: Path) -> None:
    _, library = analyze_template(reference_deck)
    # The fixture deck's slides carry no slide-number placeholders.
    assert numbered_layouts(library) == set()


def test_add_slide_number(reference_deck: Path) -> None:
    prs = Presentation(str(reference_deck))
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content has sldNum
    assert add_slide_number(slide)
    assert any(
        s.is_placeholder and s.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER
        for s in slide.shapes
    )
