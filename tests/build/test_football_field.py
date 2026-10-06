from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from deckforge.build.charts import RangeBar, axis_bounds, nice_step
from tests.conftest import build_full


def test_native_floating_bar_chart(full_inputs: Path) -> None:
    slide = Presentation(str(build_full(full_inputs))).slides[5]
    chart = next(s.chart for s in slide.shapes if s.has_chart)
    plot = chart.plots[0]
    assert [s.name for s in plot.series] == ["Low", "Range", "Pad"]
    assert list(plot.categories) == ["Comps", "DCF"]  # reversed: spec order reads top-down
    assert list(plot.series[1].values) == [7.0, 13.0]
    texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
    assert any("Current share price: $42.50" in t for t in texts)


@pytest.mark.parametrize(("span", "step"), [(131.0, 25.0), (13.0, 2.5), (900.0, 200.0)])
def test_nice_step(span: float, step: float) -> None:
    assert nice_step(span) == step


def test_axis_bounds_leave_room_and_stay_positive() -> None:
    lo, hi, step = axis_bounds([RangeBar("a", 110, 241.3, "", ""), RangeBar("b", 140, 172, "", "")])
    assert (lo, hi, step) == (50.0, 300.0, 50.0)
