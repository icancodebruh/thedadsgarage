from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from deckforge.build.renderer import load_manifest
from tests.conftest import TEST_SOURCE, build_full


def test_bullets_render_fact_tokens(full_inputs: Path) -> None:
    deck = build_full(full_inputs)
    slide = Presentation(str(deck)).slides[1]
    texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
    body = next(t for t in texts if t.startswith("Revenue"))
    assert body.split("\n") == [
        "Revenue reached $1.25bn in FY2025",
        "Testco trades at 9.5x LTM EBITDA",
    ]
    assert f"Source: {TEST_SOURCE}." in texts
    bound = {b.text for b in load_manifest(deck).for_slide(2)}
    assert bound == {"$1.25bn", "9.5x"}
