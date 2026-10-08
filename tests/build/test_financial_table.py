from __future__ import annotations

import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.table import Table

from deckforge.build.formatting import DASH
from deckforge.build.renderer import build_deck
from deckforge.build.slide_types.base import LayoutOverflowError
from deckforge.spec.models import DeckSpec
from tests.conftest import TEST_SOURCE


def _build(inputs: Path, out: Path) -> Table:
    deck = build_deck(DeckSpec.load(inputs / "deck_spec.json"), inputs, out / "deck.pptx")
    slide = Presentation(str(deck)).slides[1]
    frames = [s for s in slide.shapes if s.has_table]
    assert len(frames) == 1, "exactly one native table"
    return frames[0].table  # type: ignore[no-any-return]


def test_native_table_bound_to_facts(deck_inputs: Path, tmp_path: Path) -> None:
    table = _build(deck_inputs, tmp_path)
    grid = [[c.text for c in row.cells] for row in table.rows]
    assert grid == [
        ["", "FY 2024", "FY 2025"],
        ["Revenue", "1,000.0", "1,250.0"],
        ["% Margin", DASH, "25.5%"],
        ["Net Income", "(12.3)", "56.8"],
    ]


def test_units_and_source_footnote(deck_inputs: Path, tmp_path: Path) -> None:
    deck = build_deck(
        DeckSpec.load(deck_inputs / "deck_spec.json"), deck_inputs, tmp_path / "deck.pptx"
    )
    texts = {
        s.text_frame.text for s in Presentation(str(deck)).slides[1].shapes if s.has_text_frame
    }
    assert "($ USD in Millions)" in texts
    assert f"Source: {TEST_SOURCE}." in texts


def test_overflow_raises(deck_inputs: Path, tmp_path: Path) -> None:
    spec_path = deck_inputs / "deck_spec.json"
    spec = json.loads(spec_path.read_text())
    row = spec["slides"][1]["rows"][0]
    spec["slides"][1]["rows"] = [row] * 25
    spec_path.write_text(json.dumps(spec))
    with pytest.raises(LayoutOverflowError):
        _build(deck_inputs, tmp_path)
