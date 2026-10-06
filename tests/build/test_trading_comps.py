from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from deckforge.build.renderer import load_manifest
from tests.conftest import build_full


def test_peers_stats_then_target(full_inputs: Path) -> None:
    deck = build_full(full_inputs)
    slide = Presentation(str(deck)).slides[4]
    table = next(s.table for s in slide.shapes if s.has_table)
    grid = [[c.text for c in row.cells] for row in table.rows]
    assert grid == [
        ["Company", "EV / EBITDA"],
        ["Peer A", "10.0x"],
        ["Peer B", "12.0x"],
        ["Peer C", "17.0x"],
        ["Median", "12.0x"],
        ["Mean", "13.0x"],
        ["Testco", "9.5x"],
    ]
    derived = [b for b in load_manifest(deck).for_slide(5) if b.derived]
    assert {b.derived for b in derived} == {"median of 3 peers", "mean of 3 peers"}
    assert all(len(b.fact_ids) == 3 and b.source_refs for b in derived)
