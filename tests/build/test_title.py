from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from deckforge.build.renderer import build_deck
from deckforge.spec.models import DeckSpec


def test_title_slide(deck_inputs: Path, tmp_path: Path) -> None:
    out = build_deck(
        DeckSpec.load(deck_inputs / "deck_spec.json"), deck_inputs, tmp_path / "deck.pptx"
    )
    slide = Presentation(str(out)).slides[0]
    assert slide.shapes.title is not None
    assert slide.shapes.title.text_frame.text == "Project Testco"
    texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
    assert "Board update\nOctober 6, 2026" in texts
    assert all(t.strip() for t in texts), "no empty placeholders left behind"
