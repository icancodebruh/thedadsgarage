"""Excel <-> facts: read a `facts` sheet, or write a blank input sheet to fill in.

Use it for data SEC filings do not carry: share prices, consensus estimates, valuation
ranges. Blank values are reported as gaps rather than guessed.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from deckforge.ingest.models import Fact, Unit

SHEET = "facts"
HEADERS = ("id", "entity", "metric", "period", "value", "unit", "source_ref")


class ExcelFactsError(ValueError):
    """The workbook does not follow the facts sheet layout."""


@dataclass(frozen=True)
class ExcelFacts:
    facts: list[Fact]
    missing: list[str]  # ids whose value cell is empty


def read_facts(path: Path, sheet: str = SHEET) -> ExcelFacts:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if sheet not in workbook.sheetnames:
        raise ExcelFactsError(f"{path.name} has no '{sheet}' sheet")
    rows = workbook[sheet].iter_rows(values_only=True)
    header = tuple(str(h).strip().lower() if h else "" for h in next(rows, ()))
    if header[: len(HEADERS)] != HEADERS:
        raise ExcelFactsError(f"{path.name}!{sheet}: header must be {HEADERS}, got {header}")
    facts: list[Fact] = []
    missing: list[str] = []
    for n, row in enumerate(rows, start=2):
        fact_id, entity, metric, period, value, unit, source = (row + (None,) * 7)[:7]
        if not fact_id:
            continue
        if value is None or str(value).strip() == "":
            missing.append(str(fact_id))
            continue
        try:
            number = Decimal(str(value))
        except InvalidOperation as exc:
            raise ExcelFactsError(f"{path.name}!{sheet}!E{n}: {value!r} is not a number") from exc
        facts.append(
            Fact(
                id=str(fact_id),
                entity=str(entity),
                metric=str(metric),
                period=str(period),
                value=number,
                unit=Unit(str(unit)),
                source_ref=str(source).strip() if source else f"{path.name}!{sheet}!E{n}",
            )
        )
    workbook.close()
    return ExcelFacts(facts=facts, missing=missing)


def write_template(path: Path, rows: list[dict[str, str]]) -> Path:
    """An input sheet pre-filled with ids/labels; the user fills value and source_ref."""
    workbook = Workbook()
    ws = workbook.active
    assert ws is not None
    ws.title = SHEET
    ws.append(list(HEADERS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(h, "") for h in HEADERS])
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    for r in range(2, len(rows) + 2):
        ws.cell(row=r, column=5).fill = input_fill
        ws.cell(row=r, column=7).fill = input_fill
    units = DataValidation(type="list", formula1='"' + ",".join(u.value for u in Unit) + '"')
    ws.add_data_validation(units)
    units.add(f"F2:F{max(2, len(rows) + 1)}")
    for col, width in zip("ABCDEFG", (34, 22, 20, 10, 14, 15, 60), strict=True):
        ws.column_dimensions[col].width = width
    workbook.save(path)
    return path
