"""Helpers shared by checkers: text extraction and explicit geometry."""

from __future__ import annotations

from collections.abc import Iterator

from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.shapes.base import BaseShape
from pptx.slide import Slide

from deckforge.template.inventory import walk
from deckforge.template.ooxml import Element, first, xpath


def paragraphs(shape: BaseShape) -> Iterator[Element]:
    """Every <a:p> in a shape: shape text bodies (p:txBody) and table cells (a:txBody)."""
    yield from xpath(shape.element, ".//p:txBody/a:p | .//a:txBody/a:p")


def paragraph_text(p: Element) -> str:
    return "".join(t.text or "" for t in xpath(p, ".//a:t"))


def shape_text(shape: BaseShape) -> str:
    return "\n".join(paragraph_text(p) for p in paragraphs(shape))


def is_slide_number(shape: BaseShape) -> bool:
    return bool(
        shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER
    )


def has_own_geometry(shape: BaseShape) -> bool:
    """True when the slide sets the shape's position (not inherited from the layout)."""
    el = shape.element
    return first(el, "p:spPr/a:xfrm") is not None or first(el, "p:xfrm") is not None


def slide_texts(slide: Slide) -> Iterator[tuple[BaseShape, str]]:
    for shape in walk(slide.shapes):
        if is_slide_number(shape):
            continue
        text = shape_text(shape)
        if text.strip():
            yield shape, text


def title_of(slide: Slide) -> str | None:
    title = slide.shapes.title
    return shape_text(title).strip() if title is not None else None
