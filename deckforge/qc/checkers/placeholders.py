"""No placeholder or filler text, and no empty placeholders."""

from __future__ import annotations

import re

from pptx.shapes.autoshape import Shape

from deckforge.qc.checkers.common import is_slide_number, shape_text, slide_texts
from deckforge.qc.context import QCContext
from deckforge.qc.models import Issue, Severity

NAME = "placeholders"
FILLER = re.compile(
    r"lorem|ipsum|\bx{3,}\b|\bTBD\b|\bTBC\b|\bTODO\b|\[insert|click to (?:add|edit)|\?\?\?",
    re.IGNORECASE,
)


def check(ctx: QCContext) -> list[Issue]:
    issues: list[Issue] = []
    for n, slide in ctx.slides():
        for shape, text in slide_texts(slide):
            if m := FILLER.search(text):
                issues.append(
                    Issue(
                        checker=NAME,
                        severity=Severity.ERROR,
                        slide=n,
                        message=f"'{shape.name}' contains filler text '{m[0]}'",
                    )
                )
        for shape in slide.shapes:
            if (
                shape.is_placeholder
                and isinstance(shape, Shape)
                and not is_slide_number(shape)
                and not shape_text(shape).strip()
            ):
                issues.append(
                    Issue(
                        checker=NAME,
                        severity=Severity.ERROR,
                        slide=n,
                        message=f"empty placeholder '{shape.name}'",
                    )
                )
    return issues
