"""Text must fit its box, shapes stay on the slide and inside the template margins."""

from __future__ import annotations

import math

from lxml import etree
from pptx.shapes.base import BaseShape
from pptx.slide import Slide

from deckforge.qc.checkers.common import has_own_geometry, is_slide_number, paragraph_text
from deckforge.qc.context import QCContext
from deckforge.qc.models import Issue, Severity, SpecFix
from deckforge.template.inventory import shape_box
from deckforge.template.models import Box
from deckforge.template.ooxml import EMU_PER_INCH, Element, first, xpath
from deckforge.template.theme import TITLE_TYPES, find_matching_placeholder

NAME = "overflow"
AVG_CHAR_EM = 0.5  # average glyph width of a proportional sans font, in ems
LINE_SPACING = 1.2
FIT_TOLERANCE = 1.05
MARGIN_TOLERANCE_IN = 0.05
DEFAULT_INSET_IN = 0.1  # OOXML default left/right inset; 0.05 top/bottom


def body_chain(shape: BaseShape, slide: Slide) -> list[Element]:
    """bodyPr elements from most to least specific: shape, layout placeholder, master."""
    elements = [shape.element]
    if shape.is_placeholder:
        ph_type, idx = shape.placeholder_format.type, shape.placeholder_format.idx
        layout = slide.slide_layout
        for source, key in ((layout.placeholders, idx), (layout.slide_master.placeholders, None)):
            match = find_matching_placeholder(source, key, ph_type)
            if match is not None:
                elements.append(match.element)
    return [bp for el in elements if (bp := first(el, "p:txBody/a:bodyPr")) is not None]


def _inset(chain: list[Element], attr: str, default: float) -> float:
    value = next((bp.get(attr) for bp in chain if bp.get(attr) is not None), None)
    return int(value) / EMU_PER_INCH if value is not None else default


def autofit(chain: list[Element]) -> str:
    """'norm' (text shrinks), 'shape' (box grows) or 'none', inherited like other props."""
    for bp in chain:
        kinds = {etree.QName(c).localname for c in bp}
        if "normAutofit" in kinds:
            return "norm"
        if "spAutoFit" in kinds:
            return "shape"
        if "noAutofit" in kinds:
            return "none"
    return "none"


def text_height_in(ctx: QCContext, shape: BaseShape, slide: Slide, width_in: float) -> float:
    """Estimated rendered height of a shape's text when wrapped to `width_in`."""
    tx_body = first(shape.element, "p:txBody")
    if tx_body is None:
        return 0.0
    chain = ctx.resolver.level_chain(shape, slide.slide_layout)
    total = 0.0
    for p in xpath(tx_body, "a:p"):
        style = ctx.resolver.resolve(first(p, "a:r/a:rPr"), first(p, "a:pPr"), chain)
        size_in = style.size_pt / 72
        chars_per_line = max(1, int(width_in / (size_in * AVG_CHAR_EM)))
        lines = max(1, math.ceil(len(paragraph_text(p)) / chars_per_line))
        total += lines * size_in * LINE_SPACING
    return total


def text_width_in(ctx: QCContext, shape: BaseShape, slide: Slide) -> float:
    """Estimated width of the widest paragraph when text does not wrap."""
    tx_body = first(shape.element, "p:txBody")
    if tx_body is None:
        return 0.0
    chain = ctx.resolver.level_chain(shape, slide.slide_layout)
    widest = 0.0
    for p in xpath(tx_body, "a:p"):
        style = ctx.resolver.resolve(first(p, "a:r/a:rPr"), first(p, "a:pPr"), chain)
        widest = max(widest, len(paragraph_text(p)) * style.size_pt / 72 * AVG_CHAR_EM)
    return widest


def _fix_for(ctx: QCContext, n: int) -> SpecFix | None:
    if n > len(ctx.spec.slides):
        return None
    kind = ctx.spec.slides[n - 1].slide_type
    if kind == "financial_table":
        return SpecFix(kind="split_table", slide=n)
    if kind in ("exec_summary", "company_overview"):
        return SpecFix(kind="revise_slide", slide=n)
    return None


def check(ctx: QCContext) -> list[Issue]:
    issues: list[Issue] = []
    width, height = ctx.style.slide_size.width_in, ctx.style.slide_size.height_in
    m = ctx.style.grid.margins
    for n, slide in ctx.slides():
        for shape in slide.shapes:
            box = shape_box(shape)
            if box is None or is_slide_number(shape):
                continue
            if not has_own_geometry(shape):
                # Position comes from the layout (e.g. the title): check only that text fits.
                issues += _text_fit(ctx, shape, slide, box, n)
                continue
            if box.x < 0 or box.y < 0 or box.x + box.w > width or box.y + box.h > height:
                issues.append(
                    Issue(
                        checker=NAME,
                        severity=Severity.ERROR,
                        slide=n,
                        message=f"'{shape.name}' extends beyond the slide",
                    )
                )
            elif (
                box.x < m.left - MARGIN_TOLERANCE_IN
                or box.x + box.w > width - m.right + MARGIN_TOLERANCE_IN
            ):
                issues.append(
                    Issue(
                        checker=NAME,
                        severity=Severity.WARNING,
                        slide=n,
                        message=f"'{shape.name}' crosses the template side margins",
                    )
                )
            if getattr(shape, "has_table", False):
                needed = _table_height(ctx, shape, slide)
                limit = ctx.style.grid.content_box.y + ctx.style.grid.content_box.h
                if box.y + needed > limit + 0.01:
                    issues.append(
                        Issue(
                            checker=NAME,
                            severity=Severity.ERROR,
                            slide=n,
                            message=f"table '{shape.name}' runs past the content "
                            f"area ({box.y + needed:.2f}in > {limit:.2f}in)",
                            fix=_fix_for(ctx, n),
                        )
                    )
                continue
            issues += _text_fit(ctx, shape, slide, box, n)
    return issues


def _overlaps(a: Box, b: Box, tol: float = 0.01) -> bool:
    """Boxes intersect; a zero-height box (a rule line) counts when it crosses the other."""
    return (
        a.x < b.x + b.w - tol
        and b.x < a.x + a.w - tol
        and a.y < b.y + max(b.h, tol) - tol
        and b.y < a.y + max(a.h, tol) - tol
    )


def _obstacles(shape: BaseShape, slide: Slide) -> list[Box]:
    """Other slide shapes plus the layout's and master's own decorations."""
    layout = slide.slide_layout
    sources = [slide.shapes, layout.shapes]
    if layout.element.get("showMasterSp", "1") not in ("0", "false"):
        sources.append(layout.slide_master.shapes)
    boxes = []
    for i, shapes in enumerate(sources):
        for other in shapes:
            if other.element is shape.element or (i and other.is_placeholder):
                continue
            box = shape_box(other)
            if box is not None:
                boxes.append(box)
    return boxes


def _text_fit(ctx: QCContext, shape: BaseShape, slide: Slide, box: Box, n: int) -> list[Issue]:
    chain = body_chain(shape, slide)
    fit = autofit(chain)
    if fit == "norm":  # PowerPoint shrinks the text to fit
        return []
    inner_w = (
        box.w - _inset(chain, "lIns", DEFAULT_INSET_IN) - _inset(chain, "rIns", DEFAULT_INSET_IN)
    )
    pad_h = _inset(chain, "tIns", 0.05) + _inset(chain, "bIns", 0.05)
    if next((bp.get("wrap") for bp in chain if bp.get("wrap")), "square") == "none":
        # Unwrapped text never grows taller; it runs out of the box sideways.
        width = text_width_in(ctx, shape, slide)
        if width > inner_w * FIT_TOLERANCE:
            return [
                Issue(
                    checker=NAME,
                    severity=Severity.ERROR,
                    slide=n,
                    message=f"unwrapped text in '{shape.name}' is ~{width:.2f}in wide but the "
                    f"box is {inner_w:.2f}in",
                    fix=_title_fix(ctx, n) if _is_title(shape) else _fix_for(ctx, n),
                )
            ]
        return []
    needed = text_height_in(ctx, shape, slide, inner_w)
    if needed <= 0 or needed <= (box.h - pad_h) * FIT_TOLERANCE:
        return []
    if fit == "shape":
        # The box grows downward to fit; that is only a problem if it then hits something.
        grown = box.model_copy(update={"h": needed + pad_h})
        if not any(_overlaps(grown, other) for other in _obstacles(shape, slide)) and (
            grown.y + grown.h <= ctx.style.slide_size.height_in
        ):
            return []
        detail = f"grows to ~{grown.h:.2f}in and runs into other content"
    else:
        detail = f"needs ~{needed:.2f}in but the box is {box.h - pad_h:.2f}in high"
    return [
        Issue(
            checker=NAME,
            severity=Severity.ERROR,
            slide=n,
            message=f"text in '{shape.name}' {detail}",
            fix=_title_fix(ctx, n) if _is_title(shape) else _fix_for(ctx, n),
        )
    ]


def _is_title(shape: BaseShape) -> bool:
    return shape.is_placeholder and shape.placeholder_format.type in TITLE_TYPES


def _title_fix(ctx: QCContext, n: int) -> SpecFix | None:
    """A title that does not fit needs rewording (Claude), never a table split."""
    return SpecFix(kind="revise_slide", slide=n) if n <= len(ctx.spec.slides) else None


def _table_height(ctx: QCContext, shape: BaseShape, slide: Slide) -> float:
    """Rows grow to fit their text, so estimate each row's rendered height."""
    table = shape.table  # type: ignore[attr-defined]
    total = 0.0
    for row in table.rows:
        row_h = row.height / EMU_PER_INCH
        for c, cell in enumerate(row.cells):
            col_w = table.columns[c].width / EMU_PER_INCH
            inner_w = col_w - (cell.margin_left + cell.margin_right) / EMU_PER_INCH
            lines = 0.0
            for p in cell.text_frame.paragraphs:
                size_pt = next((r.font.size.pt for r in p.runs if r.font.size), 12.0)
                cpl = max(1, int(inner_w / (size_pt / 72 * AVG_CHAR_EM)))
                lines += max(1, math.ceil(len(p.text) / cpl)) * size_pt / 72 * LINE_SPACING
            pad = (cell.margin_top + cell.margin_bottom) / EMU_PER_INCH
            row_h = max(row_h, lines + pad)
        total += row_h
    return total
