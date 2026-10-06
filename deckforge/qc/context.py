"""Everything a checker may inspect, loaded once per QC run."""

from __future__ import annotations

from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path

import pptx
from pptx.presentation import Presentation
from pptx.shapes.base import BaseShape
from pptx.slide import Slide

from deckforge.build.manifest import Manifest
from deckforge.ingest.models import FactSet
from deckforge.spec.models import DeckSpec
from deckforge.template.inventory import walk
from deckforge.template.models import LayoutLibrary, StyleSpec
from deckforge.template.theme import TextStyleResolver, parse_theme


@dataclass
class QCContext:
    deck_path: Path
    prs: Presentation
    spec: DeckSpec
    style: StyleSpec
    library: LayoutLibrary
    facts: FactSet
    manifest: Manifest
    resolver: TextStyleResolver

    @classmethod
    def load(cls, deck: Path, spec: DeckSpec, base_dir: Path, manifest: Manifest) -> QCContext:
        prs = pptx.Presentation(str(deck))
        master = prs.slide_masters[0]
        return cls(
            deck_path=deck,
            prs=prs,
            spec=spec,
            style=StyleSpec.model_validate_json((base_dir / spec.template.style).read_text()),
            library=LayoutLibrary.model_validate_json(
                (base_dir / spec.template.layouts).read_text()
            ),
            facts=FactSet.load(base_dir / spec.facts),
            manifest=manifest,
            resolver=TextStyleResolver(prs, master, parse_theme(master)),
        )

    def slides(self) -> Iterator[tuple[int, Slide]]:
        yield from enumerate(self.prs.slides, start=1)

    def shapes(self, slide: Slide) -> Iterator[BaseShape]:
        yield from walk(slide.shapes)
