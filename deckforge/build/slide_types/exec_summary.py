"""Executive summary: one message per bullet in the layout's body placeholder."""

from __future__ import annotations

from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.slide import Slide

from deckforge.build.slide_types.base import (
    BuildContext,
    add_text,
    data_frame,
    find_placeholder,
    place,
    remove_empty_placeholders,
    set_paragraphs,
    set_title,
    source_footnote,
)
from deckforge.spec.models import ExecSummarySlide


class ExecSummarySlideType:
    slide_type = "exec_summary"

    def render(self, slide: Slide, spec: ExecSummarySlide, ctx: BuildContext) -> None:
        set_title(slide, spec.title, spec.layout_id)
        bullets = [ctx.text(b) for b in spec.bullets]
        body_box = data_frame(slide, ctx, units=None)
        body = find_placeholder(slide, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)
        if body is not None:
            # Keep the layout's bullet styling; pin geometry so the footnote zone stays clear.
            place(body, body_box)
            set_paragraphs(body.text_frame, bullets)
        else:
            add_text(slide, body_box, bullets, ctx.tokens.body)
        remove_empty_placeholders(slide)
        source_footnote(slide, ctx)
