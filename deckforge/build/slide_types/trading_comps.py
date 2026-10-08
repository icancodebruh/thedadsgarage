"""Trading comparables: peers x metrics native table with computed summary statistics."""

from __future__ import annotations

import statistics
from decimal import Decimal

from pptx.slide import Slide

from deckforge.build.formatting import DASH, units_label
from deckforge.build.slide_types.base import (
    BuildContext,
    data_frame,
    footnote_text,
    remove_empty_placeholders,
    set_title,
)
from deckforge.build.slide_types.financial_table import render_cell
from deckforge.build.tables import RowModel, draw_table
from deckforge.spec.models import (
    LITERAL_CELLS,
    CompsRow,
    FormatKind,
    RowStyle,
    Statistic,
    TradingCompsSlide,
)

LABEL_COL_SHARE = 0.26


class CompsError(ValueError):
    """A comps column mixes units across the peers used for statistics."""


def _stat(name: Statistic, values: list[Decimal]) -> Decimal:
    if name == "median":
        return Decimal(statistics.median(values))
    if name == "mean":
        return sum(values, Decimal(0)) / len(values)
    return max(values) if name == "high" else min(values)


class TradingCompsSlideType:
    slide_type = "trading_comps"

    def render(self, slide: Slide, spec: TradingCompsSlide, ctx: BuildContext) -> None:
        set_title(slide, spec.title, spec.layout_id)
        remove_empty_placeholders(slide)
        units = spec.units_label or self._units(spec, ctx)
        subtitle = ctx.text(spec.subtitle) if spec.subtitle else None
        lines = " ".join(x for x in (subtitle, units) if x) or None
        body = data_frame(slide, ctx, lines, footnote_text(ctx, spec.fact_refs(), spec.footnote))

        peers = [self._row(r, spec, ctx) for r in spec.rows if not r.is_target]
        targets = [self._row(r, spec, ctx) for r in spec.rows if r.is_target]
        rows = peers + self._stats(spec, ctx) + targets
        draw_table(
            slide,
            body,
            ["Company", *[c.header for c in spec.columns]],
            rows,
            ctx.tokens,
            label_share=LABEL_COL_SHARE,
            slide_number=ctx.slide_number,
            what=f"'{spec.title}'",
        )

    def _units(self, spec: TradingCompsSlide, ctx: BuildContext) -> str | None:
        for i, col in enumerate(spec.columns):
            if col.format.kind is FormatKind.CURRENCY and not col.format.symbol:
                ref = next(
                    (c for r in spec.rows if (c := r.cells[i]) and c not in LITERAL_CELLS), None
                )
                if ref:
                    return units_label(col.format, ctx.facts.get(ref).unit)
        return None

    def _row(self, row: CompsRow, spec: TradingCompsSlide, ctx: BuildContext) -> RowModel:
        cells = [
            render_cell(ref, col.format, ctx)
            for ref, col in zip(row.cells, spec.columns, strict=True)
        ]
        return RowModel(
            [row.company, *cells], RowStyle.HIGHLIGHT if row.is_target else RowStyle.NORMAL
        )

    def _stats(self, spec: TradingCompsSlide, ctx: BuildContext) -> list[RowModel]:
        peers = [r for r in spec.rows if not r.is_target and r.in_stats]
        out: list[RowModel] = []
        for name in spec.statistics:
            cells: list[str] = []
            for i, col in enumerate(spec.columns):
                refs = [ref for r in peers if (ref := r.cells[i]) and ref not in LITERAL_CELLS]
                if not col.stats:
                    cells.append("")
                    continue
                facts = ctx.facts.require(refs)
                if not facts:
                    cells.append(DASH)
                    continue
                units = {f.unit for f in facts}
                if len(units) > 1:
                    raise CompsError(f"column '{col.header}' mixes units {sorted(units)}")
                value = _stat(name, [f.value for f in facts])
                cells.append(
                    ctx.derived(
                        value,
                        facts[0].unit,
                        col.format,
                        [f.id for f in facts],
                        f"{name} of {len(facts)} peers",
                    )
                )
            label = f"{spec.stats_label} {name}" if spec.stats_label else name.title()
            out.append(RowModel([label, *cells], RowStyle.SUBTOTAL))
        return out
