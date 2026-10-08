"""Title slide: deck title, subtitle, and date in the layout's own placeholders."""

from __future__ import annotations

from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.slide import Slide

from deckforge.build.slide_types.base import (
    BuildContext,
    LayoutMismatchError,
    find_placeholder,
    remove_empty_placeholders,
    set_paragraphs,
    set_title,
)
from deckforge.spec.models import TitleSlide


class TitleSlideType:
    slide_type = "title"

    def render(self, slide: Slide, spec: TitleSlide, ctx: BuildContext) -> None:
        set_title(slide, spec.title, spec.layout_id)
        lines = [*(spec.subtitle or "").splitlines(), _format_date(spec) or ""]
        lines = [line for line in lines if line.strip()]
        if lines:
            sub = find_placeholder(slide, PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY)
            if sub is None:
                raise LayoutMismatchError(f"layout '{spec.layout_id}' has no subtitle placeholder")
            set_paragraphs(sub.text_frame, lines)
        remove_empty_placeholders(slide)


def _format_date(spec: TitleSlide) -> str | None:
    return f"{spec.date:%B} {spec.date.day}, {spec.date.year}" if spec.date else None
