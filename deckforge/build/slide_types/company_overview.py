"""Company overview: description + highlights (left), key facts table (right)."""

from __future__ import annotations

from pptx.enum.text import PP_ALIGN
from pptx.slide import Slide

from deckforge.build.slide_types.base import (
    GAP_IN,
    BuildContext,
    add_text,
    data_frame,
    footnote_text,
    line_height_in,
    remove_empty_placeholders,
    section_header,
    set_title,
)
from deckforge.build.tables import RowModel, draw_table
from deckforge.spec.models import CompanyOverviewSlide, KeyStat
from deckforge.template.models import Box

LEFT_SHARE = 0.5
COLUMN_GAP_IN = 0.3
NUMERIC_LABEL_SHARE = 0.6  # label column share when every value is a number
TEXT_LABEL_SHARE = 0.36  # narrower labels when some values are text


def _value(stat: KeyStat, ctx: BuildContext) -> str:
    if stat.text is not None:
        return ctx.text(stat.text)
    assert stat.fact is not None and stat.format is not None
    return ctx.number(stat.fact, stat.format)


class CompanyOverviewSlideType:
    slide_type = "company_overview"

    def render(self, slide: Slide, spec: CompanyOverviewSlide, ctx: BuildContext) -> None:
        tokens = ctx.tokens
        set_title(slide, spec.title, spec.layout_id)
        remove_empty_placeholders(slide)
        body = data_frame(slide, ctx, None, footnote_text(ctx, spec.fact_refs()))
        header_h = line_height_in(tokens.table.size_pt, 1.8)
        left_w = (body.w - COLUMN_GAP_IN) * LEFT_SHARE
        right_x = body.x + left_w + COLUMN_GAP_IN
        right_w = body.w - left_w - COLUMN_GAP_IN
        content_y = body.y + header_h + GAP_IN
        content_h = body.h - header_h - GAP_IN

        section_header(
            slide, Box(x=body.x, y=body.y, w=left_w, h=header_h), spec.left_heading, tokens
        )
        paragraphs = [ctx.text(spec.description)] + [ctx.text(h) for h in spec.highlights]
        add_text(
            slide,
            Box(x=body.x, y=content_y, w=left_w, h=content_h),
            paragraphs,
            tokens.body,
            bullets_from=1,
        )

        section_header(
            slide, Box(x=right_x, y=body.y, w=right_w, h=header_h), spec.right_heading, tokens
        )
        rows = [RowModel([s.label, _value(s, ctx)]) for s in spec.key_stats]
        share = TEXT_LABEL_SHARE if any(s.text for s in spec.key_stats) else NUMERIC_LABEL_SHARE
        draw_table(
            slide,
            Box(x=right_x, y=content_y, w=right_w, h=content_h),
            None,
            rows,
            tokens,
            label_share=share,
            slide_number=ctx.slide_number,
            what=f"'{spec.title}' key statistics",
            value_align=PP_ALIGN.LEFT if any(s.text for s in spec.key_stats) else PP_ALIGN.RIGHT,
        )
