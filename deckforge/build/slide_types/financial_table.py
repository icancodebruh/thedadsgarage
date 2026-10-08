"""Financial summary table: native table bound to facts, with units line and source."""

from __future__ import annotations

from pptx.slide import Slide

from deckforge.build.formatting import DASH, units_label
from deckforge.build.slide_types.base import (
    BuildContext,
    data_frame,
    footnote_text,
    remove_empty_placeholders,
    set_title,
)
from deckforge.build.tables import RowModel, draw_table
from deckforge.spec.models import LITERAL_CELLS, FinancialTableSlide, FormatKind, NumberFormat

LABEL_COL_SHARE = 0.26


def render_cell(ref: str | None, fmt: NumberFormat, ctx: BuildContext) -> str:
    """A fact rendered in `fmt`, a literal marker (NM, n.a.) as written, or a dash."""
    if ref is None:
        return DASH
    if ref in LITERAL_CELLS:
        return ref
    return ctx.number(ref, fmt)


class FinancialTableSlideType:
    slide_type = "financial_table"

    def render(self, slide: Slide, spec: FinancialTableSlide, ctx: BuildContext) -> None:
        set_title(slide, spec.title, spec.layout_id)
        remove_empty_placeholders(slide)
        units = spec.units_label
        if units is None:
            row = next((r for r in spec.rows if r.format.kind is FormatKind.CURRENCY), spec.rows[0])
            refs = [c for c in row.cells if c and c not in LITERAL_CELLS]
            units = units_label(row.format, ctx.facts.get(refs[0]).unit) if refs else None
        body = data_frame(slide, ctx, units, footnote_text(ctx, spec.fact_refs(), spec.footnote))
        rows = [
            RowModel([r.label] + [render_cell(ref, r.format, ctx) for ref in r.cells], r.style)
            for r in spec.rows
        ]
        draw_table(
            slide,
            body,
            ["", *spec.columns],
            rows,
            ctx.tokens,
            label_share=LABEL_COL_SHARE,
            slide_number=ctx.slide_number,
            what=f"'{spec.title}'",
        )
