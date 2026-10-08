"""Slide-type plugin interface, build context, and shared drawing helpers."""

from __future__ import annotations

import copy
import re
from collections.abc import Iterable
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any, Protocol, cast

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.slide import Slide, SlideLayout
from pptx.text.text import TextFrame, _Paragraph, _Run
from pptx.util import Emu, Inches, Pt

from deckforge.build.formatting import FACT_TOKEN, format_fact, format_value, parse_token
from deckforge.build.manifest import Binding, Manifest
from deckforge.build.style_tokens import StyleTokens, TextToken
from deckforge.ingest.models import FactSet, Unit
from deckforge.spec.models import NumberFormat
from deckforge.template.models import Box
from deckforge.template.ooxml import first, xpath


class LayoutOverflowError(ValueError):
    """Content does not fit the space the template allows."""

    def __init__(self, message: str, slide: int, max_rows: int | None = None) -> None:
        super().__init__(message)
        self.slide = slide
        self.max_rows = max_rows  # body rows that would fit, when known


class LayoutMismatchError(ValueError):
    """The chosen layout lacks a placeholder this slide type needs."""


def _fmt_key(fmt: NumberFormat) -> str:
    return f"{fmt.kind.value}/{fmt.scale.value}/{fmt.decimals}"


@dataclass
class BuildContext:
    """Shared state for one build. Every number goes through `number`/`derived`/`text`."""

    tokens: StyleTokens
    facts: FactSet
    manifest: Manifest = field(default_factory=Manifest)
    slide_number: int = 0

    def number(self, fact_id: str, fmt: NumberFormat, *, inline: bool = False) -> str:
        fact = self.facts.get(fact_id)
        text = format_fact(fact, fmt, inline=inline)
        self.manifest.bindings.append(
            Binding(
                slide=self.slide_number,
                text=text,
                fact_ids=[fact.id],
                source_refs=[fact.source_ref],
                format_key=_fmt_key(fmt),
            )
        )
        return text

    def derived(
        self, value: Decimal, unit: Unit, fmt: NumberFormat, fact_ids: list[str], how: str
    ) -> str:
        """A value computed from facts (e.g. a peer median); sources are the inputs'."""
        text = format_value(value, unit, fmt)
        sources = sorted({f.source_ref for f in self.facts.require(fact_ids)})
        self.manifest.bindings.append(
            Binding(
                slide=self.slide_number,
                text=text,
                fact_ids=fact_ids,
                source_refs=sources,
                derived=how,
                format_key=_fmt_key(fmt),
            )
        )
        return text

    def text(self, template: str) -> str:
        """Render `{{fact|fmt}}` tokens in running text."""

        def sub(match: Any) -> str:
            fact_id, fmt = parse_token(match)
            return self.number(fact_id, fmt, inline=True)

        return FACT_TOKEN.sub(sub, template)

    def sources_on_slide(self) -> list[str]:
        return sorted(
            {s for b in self.manifest.for_slide(self.slide_number) for s in b.source_refs}
        )


class SlideType(Protocol):
    """A plugin that renders one spec entry onto a slide created from its layout."""

    slide_type: str

    def render(self, slide: Slide, spec: Any, ctx: BuildContext) -> None: ...


# --- placeholders -----------------------------------------------------------------------


def placeholders(slide: Slide | SlideLayout) -> list[BaseShape]:
    # python-pptx annotates placeholder iteration incorrectly; cast through Iterable.
    return list(cast("Iterable[BaseShape]", slide.placeholders))


def find_placeholder(slide: Slide, *types: PP_PLACEHOLDER) -> Shape | None:
    for ph in placeholders(slide):
        if ph.placeholder_format.type in types and isinstance(ph, Shape):
            return ph
    return None


def set_title(slide: Slide, title: str, layout_id: str) -> None:
    ph = find_placeholder(slide, PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
    if ph is None:
        raise LayoutMismatchError(f"layout '{layout_id}' has no title placeholder")
    ph.text_frame.text = title


def add_slide_number(slide: Slide) -> bool:
    """Copy the layout's slide-number placeholder (with its field) onto the slide.

    python-pptx's `add_slide` skips date/footer/slide-number placeholders.
    Returns False when the layout has none.
    """
    layout_ph = next(
        (
            ph
            for ph in placeholders(slide.slide_layout)
            if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER
        ),
        None,
    )
    if layout_ph is None:
        return False
    sp = copy.deepcopy(layout_ph.element)
    tree = first(slide.element, "p:cSld/p:spTree")
    ids = [int(el.get("id")) for el in xpath(tree, ".//p:cNvPr")]
    c_nv_pr = first(sp, "p:nvSpPr/p:cNvPr")
    assert tree is not None and c_nv_pr is not None
    c_nv_pr.set("id", str(max(ids, default=1) + 1))
    tree.append(sp)
    return True


def remove_shape(shape: BaseShape) -> None:
    el = shape.element
    el.getparent().remove(el)


def remove_empty_placeholders(slide: Slide) -> None:
    """Delete placeholders left without text so no prompt text renders."""
    for ph in placeholders(slide):
        if not (isinstance(ph, Shape) and ph.text_frame.text.strip()):
            remove_shape(ph)


def place(shape: BaseShape, box: Box) -> None:
    shape.left, shape.top = Inches(box.x), Inches(box.y)
    shape.width, shape.height = Inches(box.w), Inches(box.h)


# --- text -------------------------------------------------------------------------------


def line_height_in(size_pt: float, factor: float = 1.4) -> float:
    return size_pt * factor / 72


def style_run(
    run: _Run, token: TextToken, *, bold: bool | None = None, color: str | None = None
) -> None:
    run.font.name = token.font
    run.font.size = Pt(token.size_pt)
    run.font.bold = token.bold if bold is None else bold
    run.font.color.rgb = RGBColor.from_string(color or token.color)  # type: ignore[no-untyped-call]


def set_paragraphs(frame: TextFrame, lines: list[str]) -> None:
    """Replace text keeping inherited (placeholder) formatting: one paragraph per line."""
    frame.text = lines[0]
    for line in lines[1:]:
        frame.add_paragraph().text = line  # type: ignore[no-untyped-call]


def _zero_margins(frame: TextFrame) -> None:
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(frame, side, Emu(0))


BULLET_INDENT_IN = 0.18


def set_bullet(paragraph: _Paragraph, char: str = "\u2022") -> None:
    """Give a paragraph a real hanging bullet (never a typed bullet character)."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("marL", str(Inches(BULLET_INDENT_IN)))
    p_pr.set("indent", str(-Inches(BULLET_INDENT_IN)))
    bullet = etree.SubElement(p_pr, qn("a:buChar"))
    bullet.set("char", char)


def add_text(
    slide: Slide,
    box: Box,
    text: str | list[str],
    token: TextToken,
    *,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    bold: bool | None = None,
    color: str | None = None,
    fill: str | None = None,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
    pad_in: float = 0.0,
    bullets_from: int | None = None,
) -> Shape:
    """A text box styled from tokens; a list renders one paragraph per item."""
    shape = slide.shapes.add_textbox(Inches(box.x), Inches(box.y), Inches(box.w), Inches(box.h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    _zero_margins(frame)
    if pad_in:
        frame.margin_left = frame.margin_right = Inches(pad_in)
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)  # type: ignore[no-untyped-call]
    lines = [text] if isinstance(text, str) else text
    for i, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = align
        if bullets_from is not None and i >= bullets_from:
            set_bullet(paragraph)
        if i:
            paragraph.space_before = Pt(token.size_pt * 0.5)
        run = paragraph.add_run()
        run.text = line
        style_run(run, token, bold=bold, color=color)
    return shape


def section_header(slide: Slide, box: Box, text: str, tokens: StyleTokens) -> Shape:
    """A filled header bar (template header fill + contrasting text)."""
    return add_text(
        slide,
        box,
        text,
        tokens.table,
        bold=True,
        color=tokens.header_text,
        fill=tokens.header_fill,
        anchor=MSO_ANCHOR.MIDDLE,
        pad_in=0.08,
    )


# --- data-slide frame -------------------------------------------------------------------

GAP_IN = 0.1


FOOTNOTE_LINE = 1.3  # footnote line height as a multiple of its font size
AVG_CHAR_EM = 0.5  # average glyph width, used to estimate how many lines text wraps to


COMPILED_FROM = re.compile(r"^(?P<body>.*?)\s*\[(?P<doc>[^\]]+?), p\.(?P<page>\d+)\]$")


def citations(source_refs: list[str]) -> tuple[list[str], dict[str, set[int]]]:
    """Split source refs into unique citations, plus pages of any compiled document.

    A ref may end with "[<document>, p.<n>]" when the fact was taken from a compiled
    document (an existing deck) that itself cites the underlying source.
    """
    items: list[str] = []
    pages: dict[str, set[int]] = {}
    for ref in source_refs:
        m = COMPILED_FROM.match(ref)
        body = ref
        if m:
            body = m["body"]
            pages.setdefault(m["doc"], set()).add(int(m["page"]))
        parts: list[str] = []
        for part in (x.strip() for x in body.split("; ")):
            if parts and part[:1].isdigit():  # "...: 6.29m shares; 49.69m post-issue"
                parts[-1] += f", {part}"
            elif part:
                parts.append(part)
        items += [x.rstrip(".") for x in parts if x.rstrip(".") not in items]
    # Drop a citation that a fuller one already covers ("QIP filing" vs "QIP filing: ...").
    items = [x for x in items if not any(o != x and o.startswith(x) for o in items)]
    return items, pages


def footnote_text(ctx: BuildContext, refs: list[str], note: str | None = None) -> str:
    """'Source: ...' for every fact the slide uses, plus an optional methodology note.

    Sources come from the slide's fact references up front, so the footnote can be
    sized before the body is laid out.
    """
    parts: list[str] = []
    if refs:
        items, pages = citations([f.source_ref for f in ctx.facts.require(refs)])
        parts.append("Source: " + "; ".join(sorted(items)) + ".")
        for doc, nums in sorted(pages.items()):
            listed = ", ".join(str(n) for n in sorted(nums))
            parts.append(f"Compiled from the {doc} (p. {listed}).")
    if note:
        parts.append("Note: " + ctx.text(note))
    return " ".join(parts)


def footnote_height(ctx: BuildContext, text: str) -> float:
    if not text:
        return 0.0
    size_in = ctx.tokens.footnote.size_pt / 72
    chars_per_line = max(1, int(ctx.tokens.content_box.w / (size_in * AVG_CHAR_EM)))
    lines = -(-len(text) // chars_per_line)
    return lines * size_in * FOOTNOTE_LINE + 0.02


def data_frame(slide: Slide, ctx: BuildContext, units: str | None, footnote: str = "") -> Box:
    """Draw the units line (top) and footnote (bottom); return the body box between them."""
    cb = ctx.tokens.content_box
    top = cb.y
    if units:
        line_h = line_height_in(ctx.tokens.footnote.size_pt)
        add_text(slide, Box(x=cb.x, y=cb.y, w=cb.w, h=line_h), units, ctx.tokens.footnote)
        top += line_h + GAP_IN
    bottom = cb.y + cb.h
    if footnote:
        h = footnote_height(ctx, footnote)
        add_text(
            slide,
            Box(x=cb.x, y=round(bottom - h, 3), w=cb.w, h=round(h, 3)),
            footnote,
            ctx.tokens.footnote,
        )
        bottom -= h + GAP_IN
    return Box(x=cb.x, y=round(top, 3), w=cb.w, h=round(bottom - top, 3))
