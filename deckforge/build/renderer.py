"""Renders a validated deck spec into the reference template's layouts."""

from __future__ import annotations

import hashlib
import logging
import zipfile
from pathlib import Path

import pptx
from pptx.presentation import Presentation
from pptx.slide import SlideLayout

from deckforge.build.slide_types import REGISTRY
from deckforge.build.slide_types.base import BuildContext, add_slide_number
from deckforge.build.style_tokens import derive_tokens
from deckforge.ingest.models import FactSet
from deckforge.spec.models import DeckSpec, FinancialTableSlide
from deckforge.template.layouts import layout_ids
from deckforge.template.models import LayoutLibrary, StyleSpec

log = logging.getLogger(__name__)

FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)  # deterministic archive timestamps


class SpecError(ValueError):
    """The spec does not fit its template, style, or layouts."""


def _clear_slides(prs: Presentation) -> None:
    """Remove the reference deck's own slides; masters and layouts stay."""
    # python-pptx has no public slide-delete API; dropping the relationship makes the
    # slide part (and its notes/media) unreachable, so it is not written on save.
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        prs.part.drop_rel(sld_id.rId)
        sld_id_lst.remove(sld_id)


def _layouts_by_id(prs: Presentation) -> dict[str, SlideLayout]:
    ids = layout_ids(prs)
    return {
        ids[id(layout.element)]: layout
        for master in prs.slide_masters
        for layout in master.slide_layouts
    }


def _normalize_zip(path: Path) -> None:
    """Rewrite the package with fixed timestamps so identical inputs give identical bytes."""
    with zipfile.ZipFile(path) as src:
        entries = [(info.filename, src.read(info)) for info in src.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for name, data in entries:
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, data)


def numbered_layouts(library: LayoutLibrary) -> set[str]:
    """Layouts whose reference slides mostly show a slide number (learned, not assumed)."""
    shown: dict[str, list[bool]] = {}
    for s in library.slides:
        has = any(sh.placeholder_type == "slide_number" for sh in s.shapes)
        shown.setdefault(s.layout_id, []).append(has)
    return {lid for lid, flags in shown.items() if sum(flags) > len(flags) / 2}


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def build_deck(spec: DeckSpec, base_dir: Path, out_path: Path) -> Path:
    template = base_dir / spec.template.pptx
    style = StyleSpec.model_validate_json((base_dir / spec.template.style).read_text("utf-8"))
    library = LayoutLibrary.model_validate_json(
        (base_dir / spec.template.layouts).read_text("utf-8")
    )
    if style.source.sha256 != _sha256(template):
        raise SpecError(f"style.json was not generated from {template.name}; re-run analyze")
    facts = FactSet.load(base_dir / spec.facts)

    known = {lay.layout_id for lay in library.layouts}
    bad = sorted({s.layout_id for s in spec.slides} - known)
    if bad:
        raise SpecError(f"unknown layout_id(s) {bad}; layouts.json has {sorted(known)}")
    # Report every missing number before building anything.
    facts.require(
        [r for s in spec.slides if isinstance(s, FinancialTableSlide) for r in s.fact_refs()]
    )

    prs = pptx.Presentation(str(template))
    _clear_slides(prs)
    layouts = _layouts_by_id(prs)
    ctx = BuildContext(tokens=derive_tokens(style), facts=facts)
    numbered = numbered_layouts(library)
    for n, slide_spec in enumerate(spec.slides, start=1):
        slide = prs.slides.add_slide(layouts[slide_spec.layout_id])
        REGISTRY[slide_spec.slide_type].render(slide, slide_spec, ctx)
        if slide_spec.layout_id in numbered:
            add_slide_number(slide)
        log.debug("slide built", extra={"slide": n, "slide_type": slide_spec.slide_type})

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    _normalize_zip(out_path)
    log.info("deck built", extra={"path": str(out_path), "slides": len(spec.slides)})
    return out_path
