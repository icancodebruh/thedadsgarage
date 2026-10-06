"""Native PowerPoint table drawing shared by every tabular slide type."""

from __future__ import annotations

from dataclasses import dataclass

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.shapes.graphfrm import GraphicFrame
from pptx.slide import Slide
from pptx.table import _Cell
from pptx.util import Emu, Inches

from deckforge.build.slide_types.base import LayoutOverflowError, style_run
from deckforge.build.style_tokens import StyleTokens, TextToken
from deckforge.spec.models import RowStyle
from deckforge.template.models import Box

ROW_HEIGHT = 1.8  # row height as a multiple of font size
CELL_PAD_X_IN, CELL_PAD_Y_IN = 0.05, 0.02
MEMO_INDENT_IN = 0.2


@dataclass(frozen=True)
class RowModel:
    cells: list[str]  # first cell is the row label
    style: RowStyle = RowStyle.NORMAL


def row_height_in(tokens: StyleTokens) -> float:
    return tokens.table.size_pt * ROW_HEIGHT / 72


def table_height_in(n_rows: int, tokens: StyleTokens) -> float:
    return row_height_in(tokens) * n_rows


def draw_table(
    slide: Slide,
    box: Box,
    header: list[str] | None,
    rows: list[RowModel],
    tokens: StyleTokens,
    *,
    label_share: float,
    slide_number: int,
    what: str,
) -> GraphicFrame:
    """Draw a table at the top of `box`; raise if it cannot fit its height."""
    n_rows = len(rows) + (1 if header else 0)
    height = table_height_in(n_rows, tokens)
    if height > box.h + 1e-6:
        capacity = int(box.h // row_height_in(tokens)) - (1 if header else 0)
        raise LayoutOverflowError(
            f"{what}: {n_rows} rows need {height:.2f}in, only {box.h:.2f}in available",
            slide_number,
            max_rows=max(capacity, 1),
        )
    n_cols = len(header) if header else len(rows[0].cells)
    frame = slide.shapes.add_table(
        n_rows, n_cols, Inches(box.x), Inches(box.y), Inches(box.w), Inches(height)
    )
    table = frame.table
    table.first_row = header is not None
    table.horz_banding = False

    total_w = Inches(box.w)
    label_w = int(total_w * label_share) if n_cols > 1 else total_w
    value_w = (total_w - label_w) // max(1, n_cols - 1)
    table.columns[0].width = Emu(label_w)
    for i in range(1, n_cols):
        table.columns[i].width = Emu(value_w)
    if n_cols > 1:
        table.columns[n_cols - 1].width = Emu(total_w - label_w - value_w * (n_cols - 2))
    for row in table.rows:
        row.height = Emu(Inches(height) // n_rows)

    offset = 0
    if header:
        for c, text in enumerate(header):
            _cell(
                table.cell(0, c),
                text,
                tokens.table,
                bold=True,
                fill=tokens.header_fill,
                color=tokens.header_text,
                align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT,
            )
        offset = 1
    for r, row in enumerate(rows, start=offset):
        bold = row.style in (RowStyle.TOTAL, RowStyle.SUBTOTAL, RowStyle.HIGHLIGHT)
        fill = tokens.emphasis_fill if row.style in (RowStyle.TOTAL, RowStyle.HIGHLIGHT) else None
        for c, text in enumerate(row.cells):
            indent = MEMO_INDENT_IN if (c == 0 and row.style is RowStyle.MEMO) else 0.0
            _cell(
                table.cell(r, c),
                text,
                tokens.table,
                bold=bold,
                fill=fill,
                indent=indent,
                align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT,
            )
    return frame


def _cell(
    cell: _Cell,
    text: str,
    token: TextToken,
    *,
    bold: bool,
    align: PP_ALIGN,
    fill: str | None,
    color: str | None = None,
    indent: float = 0.0,
) -> None:
    if fill:
        cell.fill.solid()  # type: ignore[no-untyped-call]
        cell.fill.fore_color.rgb = RGBColor.from_string(fill)  # type: ignore[no-untyped-call]
    else:
        cell.fill.background()  # type: ignore[no-untyped-call]  # explicit no-fill beats the table style
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(CELL_PAD_X_IN + indent)
    cell.margin_right = Inches(CELL_PAD_X_IN)
    cell.margin_top = cell.margin_bottom = Inches(CELL_PAD_Y_IN)
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    style_run(run, token, bold=bold, color=color)
