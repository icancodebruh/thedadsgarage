"""Walks slides and records observed text styles, colors, and shape geometry."""

from __future__ import annotations

from collections.abc import Iterator
from dataclasses import dataclass, field

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.slide import Slide

from deckforge.template.models import Box, ShapeKind, TextRole
from deckforge.template.ooxml import Element, emu_to_in, first, xpath
from deckforge.template.theme import TITLE_TYPES, TextStyleResolver

FILL_TAGS = ("a:solidFill", "a:noFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill")


@dataclass(frozen=True)
class TextObs:
    role: TextRole
    font: str
    size_pt: float
    bold: bool
    color: str | None
    chars: int


@dataclass
class ColorObs:
    text: dict[str, int] = field(default_factory=dict)  # weighted by characters
    fill: dict[str, int] = field(default_factory=dict)  # weighted by occurrences
    line: dict[str, int] = field(default_factory=dict)

    @staticmethod
    def add(bucket: dict[str, int], hex_: str | None, weight: int = 1) -> None:
        if hex_:
            bucket[hex_] = bucket.get(hex_, 0) + weight


def walk(shapes: object) -> Iterator[BaseShape]:
    """Yield every shape, descending into groups."""
    for shape in shapes:  # type: ignore[attr-defined]
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(shape.shapes)


def shape_box(shape: BaseShape) -> Box | None:
    if shape.left is None or shape.top is None or shape.width is None or shape.height is None:
        return None
    return Box(
        x=emu_to_in(shape.left),
        y=emu_to_in(shape.top),
        w=emu_to_in(shape.width),
        h=emu_to_in(shape.height),
    )


def shape_kind(shape: BaseShape) -> ShapeKind:
    if shape.is_placeholder:
        return ShapeKind.PLACEHOLDER
    st = shape.shape_type
    mapping = {
        MSO_SHAPE_TYPE.TEXT_BOX: ShapeKind.TEXT,
        MSO_SHAPE_TYPE.AUTO_SHAPE: ShapeKind.AUTOSHAPE,
        MSO_SHAPE_TYPE.PICTURE: ShapeKind.PICTURE,
        MSO_SHAPE_TYPE.TABLE: ShapeKind.TABLE,
        MSO_SHAPE_TYPE.CHART: ShapeKind.CHART,
        MSO_SHAPE_TYPE.GROUP: ShapeKind.GROUP,
        MSO_SHAPE_TYPE.FREEFORM: ShapeKind.FREEFORM,
        MSO_SHAPE_TYPE.LINE: ShapeKind.LINE,
    }
    if st in mapping:
        return mapping[st]
    if getattr(shape, "has_chart", False):
        return ShapeKind.CHART
    if getattr(shape, "has_table", False):
        return ShapeKind.TABLE
    return ShapeKind.OTHER


def placeholder_type(shape: BaseShape) -> PP_PLACEHOLDER | None:
    return shape.placeholder_format.type if shape.is_placeholder else None


def shape_text(shape: BaseShape) -> str:
    """Stripped text of a shape with a text frame, else an empty string."""
    return shape.text_frame.text.strip() if isinstance(shape, Shape) else ""


def text_role(shape: BaseShape) -> TextRole:
    ph = placeholder_type(shape)
    if ph in TITLE_TYPES:
        return TextRole.TITLE
    if ph is not None and ph not in (
        PP_PLACEHOLDER.DATE,
        PP_PLACEHOLDER.FOOTER,
        PP_PLACEHOLDER.SLIDE_NUMBER,
    ):
        return TextRole.BODY
    return TextRole.OTHER


def _runs(
    tx_body: Element, resolver: TextStyleResolver, chain: list[Element], role: TextRole
) -> Iterator[TextObs]:
    for p in xpath(tx_body, "a:p"):
        p_pr = first(p, "a:pPr")
        for r in xpath(p, "a:r"):
            t = first(r, "a:t")
            text = (t.text or "").strip() if t is not None else ""
            if not text:
                continue
            s = resolver.resolve(first(r, "a:rPr"), p_pr, chain)
            yield TextObs(role, s.font, s.size_pt, s.bold, s.color, len(text))


def _shape_fill_and_line(sp: Element, resolver: TextStyleResolver, colors: ColorObs) -> None:
    sp_pr = first(sp, "p:spPr")
    if sp_pr is None:
        return
    has_explicit_fill = any(first(sp_pr, tag) is not None for tag in FILL_TAGS)
    if has_explicit_fill:
        ColorObs.add(colors.fill, resolver.colors.solid_fill(sp_pr))
    else:  # fall back to the shape's style reference (<p:style><a:fillRef>)
        ref = first(sp, "p:style/a:fillRef/*")
        if ref is not None and first(sp, "p:style/a:fillRef[@idx!='0']") is not None:
            ColorObs.add(colors.fill, resolver.colors.resolve(ref))
    ln = first(sp_pr, "a:ln")
    if ln is not None:
        if first(ln, "a:noFill") is None:
            ColorObs.add(colors.line, resolver.colors.solid_fill(ln))
    else:
        ref = first(sp, "p:style/a:lnRef[@idx!='0']/*")
        ColorObs.add(colors.line, resolver.colors.resolve(ref))


def inventory_slide(slide: Slide, resolver: TextStyleResolver, colors: ColorObs) -> list[TextObs]:
    """Collect text observations and accumulate colors for one slide."""
    layout = slide.slide_layout
    texts: list[TextObs] = []
    for shape in walk(slide.shapes):
        el = shape.element
        if shape.shape_type in (
            MSO_SHAPE_TYPE.AUTO_SHAPE,
            MSO_SHAPE_TYPE.TEXT_BOX,
            MSO_SHAPE_TYPE.FREEFORM,
            MSO_SHAPE_TYPE.PLACEHOLDER,
            MSO_SHAPE_TYPE.LINE,
        ):
            _shape_fill_and_line(el, resolver, colors)
        tx_body = first(el, "p:txBody")
        if tx_body is not None:
            texts.extend(
                _runs(tx_body, resolver, resolver.level_chain(shape, layout), text_role(shape))
            )
        if getattr(shape, "has_table", False):
            chain = resolver.level_chain(None, None)
            for row_i, tr in enumerate(xpath(el, ".//a:tbl/a:tr")):
                role = TextRole.TABLE_HEADER if row_i == 0 else TextRole.TABLE_BODY
                for tc in xpath(tr, "a:tc"):
                    ColorObs.add(colors.fill, resolver.colors.solid_fill(first(tc, "a:tcPr")))
                    cell_body = first(tc, "a:txBody")
                    if cell_body is not None:
                        texts.extend(_runs(cell_body, resolver, chain, role))
    for obs in texts:
        ColorObs.add(colors.text, obs.color, obs.chars)
    return texts
