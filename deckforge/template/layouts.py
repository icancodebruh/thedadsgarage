"""Builds the layout library: master layouts plus per-slide reference archetypes."""

from __future__ import annotations

import re
from collections import Counter
from collections.abc import Iterable
from typing import cast

from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.presentation import Presentation
from pptx.shapes.base import BaseShape
from pptx.slide import Slide, SlideLayout

from deckforge.template.inventory import placeholder_type, shape_box, shape_kind, shape_text
from deckforge.template.models import (
    Box,
    ContentKind,
    LayoutSpec,
    PlaceholderSpec,
    ShapeKind,
    SlideArchetype,
    SlideShape,
    StaticShape,
)
from deckforge.template.theme import OTHER_TYPES, TITLE_TYPES

PICTURE_AS_DATA_RATIO = 0.15  # a picture this large on a content slide is likely data
FOOTER_TOP_RATIO = 0.88
FOOTER_MAX_H_IN = 0.6

_KIND_BUCKET: dict[ShapeKind, ContentKind] = {
    ShapeKind.TABLE: ContentKind.TABLE,
    ShapeKind.CHART: ContentKind.CHART,
    ShapeKind.PICTURE: ContentKind.PICTURE,
    ShapeKind.TEXT: ContentKind.TEXT,
    ShapeKind.PLACEHOLDER: ContentKind.TEXT,
    ShapeKind.AUTOSHAPE: ContentKind.DIAGRAM,
    ShapeKind.FREEFORM: ContentKind.DIAGRAM,
    ShapeKind.GROUP: ContentKind.DIAGRAM,
    ShapeKind.LINE: ContentKind.DIAGRAM,
}


def slugify(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_") or "layout"


def layout_ids(prs: Presentation) -> dict[int, str]:
    """Stable, unique id per layout keyed by the layout's `id()`; deduped in deck order."""
    ids: dict[int, str] = {}
    seen: Counter[str] = Counter()
    for m_i, master in enumerate(prs.slide_masters):
        for layout in master.slide_layouts:
            slug = slugify(layout.name)
            seen[slug] += 1
            ids[id(layout.element)] = slug if seen[slug] == 1 else f"{slug}_{m_i}_{seen[slug]}"
    return ids


def is_footer(shape: BaseShape, box: Box, slide_h: float) -> bool:
    if placeholder_type(shape) in OTHER_TYPES:
        return True
    return box.y >= FOOTER_TOP_RATIO * slide_h and box.h <= FOOTER_MAX_H_IN


def _static(shapes: Iterable[BaseShape]) -> list[StaticShape]:
    out: list[StaticShape] = []
    for shape in shapes:
        box = shape_box(shape)
        if shape.is_placeholder or box is None:
            continue
        out.append(
            StaticShape(
                kind=shape_kind(shape).value,
                name=shape.name,
                box=box,
                text=shape_text(shape) or None,
            )
        )
    return out


def build_layouts(prs: Presentation, ids: dict[int, str]) -> list[LayoutSpec]:
    usage: dict[str, list[int]] = {}
    for n, slide in enumerate(prs.slides, start=1):
        usage.setdefault(ids[id(slide.slide_layout.element)], []).append(n)

    specs: list[LayoutSpec] = []
    for m_i, master in enumerate(prs.slide_masters):
        master_static = _static(master.shapes)
        for l_i, layout in enumerate(master.slide_layouts):
            lid = ids[id(layout.element)]
            shows_master = layout.element.get("showMasterSp", "1") not in ("0", "false")
            placeholders = [
                PlaceholderSpec(
                    idx=ph.placeholder_format.idx,
                    type=_ph_name(ph.placeholder_format.type),
                    name=ph.name,
                    box=shape_box(ph),
                )
                for ph in layout.placeholders
            ]
            specs.append(
                LayoutSpec(
                    layout_id=lid,
                    name=layout.name,
                    master_index=m_i,
                    layout_index=l_i,
                    placeholders=sorted(placeholders, key=lambda p: p.idx),
                    static_shapes=(master_static if shows_master else []) + _static(layout.shapes),
                    used_by_slides=usage.get(lid, []),
                )
            )
    return specs


def _ph_name(ph_type: PP_PLACEHOLDER | None) -> str:
    return ph_type.name.lower() if ph_type is not None else "unknown"


def _is_title_layout(layout: SlideLayout) -> bool:
    # python-pptx annotates placeholder iteration incorrectly; cast through Iterable.
    for ph in cast("Iterable[BaseShape]", layout.placeholders):
        if ph.placeholder_format.type == PP_PLACEHOLDER.CENTER_TITLE:
            return True
    return False


def build_archetype(
    slide: Slide, number: int, layout_id: str, width: float, height: float
) -> SlideArchetype:
    shapes: list[SlideShape] = []
    area: Counter[ContentKind] = Counter()
    warnings: list[str] = []
    slide_area = width * height
    for shape in slide.shapes:
        box = shape_box(shape)
        if box is None:
            continue
        kind = shape_kind(shape)
        ph = placeholder_type(shape)
        shapes.append(
            SlideShape(
                kind=kind,
                name=shape.name,
                box=box,
                placeholder_type=_ph_name(ph) if shape.is_placeholder else None,
                has_text=bool(shape_text(shape)),
            )
        )
        if (
            box.x < -0.01
            or box.y < -0.01
            or box.x + box.w > width + 0.01
            or (box.y + box.h > height + 0.01)
        ):
            warnings.append(f"'{shape.name}' extends beyond the slide bounds")
        if ph in TITLE_TYPES or is_footer(shape, box, height):
            continue
        share = box.w * box.h / slide_area
        area[_KIND_BUCKET.get(kind, ContentKind.DIAGRAM)] += round(share * 10_000)
        if kind is ShapeKind.PICTURE and share >= PICTURE_AS_DATA_RATIO:
            warnings.append(
                f"'{shape.name}' is a picture covering {share:.0%} of the slide; likely a "
                "pasted table/chart - rebuild as a native object"
            )

    if _is_title_layout(slide.slide_layout):
        content_kind = ContentKind.TITLE
    elif area:
        content_kind = max(area.items(), key=lambda kv: (kv[1], kv[0].value))[0]
    else:
        content_kind = ContentKind.EMPTY
    title_shape = slide.shapes.title
    title = title_shape.text_frame.text.strip() if title_shape is not None else None
    return SlideArchetype(
        slide_number=number,
        layout_id=layout_id,
        title=title or None,
        content_kind=content_kind,
        shape_counts=dict(sorted(Counter(s.kind.value for s in shapes).items())),
        shapes=shapes,
        warnings=warnings,
    )
