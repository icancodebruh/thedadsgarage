"""Template analysis entry point: reference .pptx -> `style.json` + `layouts.json`."""

from __future__ import annotations

import hashlib
import logging
from collections import defaultdict
from collections.abc import Iterable, Mapping
from pathlib import Path
from typing import TypeVar

from pptx import Presentation

from deckforge.template.grid import Placed, Zone, infer_grid
from deckforge.template.inventory import ColorObs, TextObs, inventory_slide, shape_box
from deckforge.template.layouts import build_archetype, build_layouts, is_footer, layout_ids
from deckforge.template.models import (
    Aspect,
    ColorUsage,
    FontUsage,
    LayoutLibrary,
    MasterTextStyles,
    Palette,
    SlideSize,
    SourceInfo,
    StyleSpec,
    TextRole,
)
from deckforge.template.text_rules import infer_text_rules
from deckforge.template.theme import THEME_SLOTS, TITLE_TYPES, TextStyleResolver, parse_theme

log = logging.getLogger(__name__)

ALLOWED_SHARE_MIN = 0.01  # a font/size must carry >=1% of characters to count as "allowed"
OFFICE_2007_THEME = {
    "dk2": "1F497D", "lt2": "EEECE1", "accent1": "4F81BD", "accent2": "C0504D",
    "accent3": "9BBB59", "accent4": "8064A2", "accent5": "4BACC6", "accent6": "F79646",
}  # fmt: skip
_ROLE_ORDER = list(TextRole)
T = TypeVar("T", str, float)
_ASPECTS: dict[Aspect, float] = {"16:9": 16 / 9, "4:3": 4 / 3, "16:10": 16 / 10}


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def _aspect(w: float, h: float) -> SlideSize:
    ratio = w / h
    aspect: Aspect = "other"
    for name, r in _ASPECTS.items():
        if abs(ratio - r) < 0.01:
            aspect = name
    return SlideSize(width_in=round(w, 3), height_in=round(h, 3), aspect=aspect)


def aggregate_fonts(obs: Iterable[TextObs]) -> list[FontUsage]:
    totals: dict[tuple[TextRole, str, float, bool, str | None], int] = defaultdict(int)
    role_chars: dict[TextRole, int] = defaultdict(int)
    for o in obs:
        totals[(o.role, o.font, o.size_pt, o.bold, o.color)] += o.chars
        role_chars[o.role] += o.chars
    usages = [
        FontUsage(
            role=role,
            font=font,
            size_pt=size,
            bold=bold,
            color=color,
            chars=chars,
            share=round(chars / role_chars[role], 4),
        )
        for (role, font, size, bold, color), chars in totals.items()
    ]
    return sorted(
        usages,
        key=lambda u: (_ROLE_ORDER.index(u.role), -u.chars, u.font, u.size_pt, u.color or ""),
    )


def allowed_values(obs: list[TextObs], key: str) -> list[T]:
    total = sum(o.chars for o in obs) or 1
    counts: dict[T, int] = defaultdict(int)
    for o in obs:
        counts[getattr(o, key)] += o.chars
    return sorted(v for v, c in counts.items() if c / total >= ALLOWED_SHARE_MIN)


def color_usages(bucket: Mapping[str, int], theme: Mapping[str, str]) -> list[ColorUsage]:
    total = sum(bucket.values()) or 1
    by_hex = {hex_: slot for slot in reversed(THEME_SLOTS) if (hex_ := theme.get(slot))}
    return [
        ColorUsage(hex=h, count=c, share=round(c / total, 4), theme_slot=by_hex.get(h))
        for h, c in sorted(bucket.items(), key=lambda kv: (-kv[1], kv[0]))
    ]


def analyze_template(path: Path) -> tuple[StyleSpec, LayoutLibrary]:
    prs = Presentation(str(path))
    if prs.slide_width is None or prs.slide_height is None:
        raise ValueError(f"{path} has no slide size")
    width, height = prs.slide_width / 914_400, prs.slide_height / 914_400
    source = SourceInfo(filename=path.name, sha256=_sha256(path), slide_count=len(prs.slides))
    masters = list(prs.slide_masters)
    master = masters[0]
    theme = parse_theme(master)
    resolver = TextStyleResolver(prs, master, theme)
    ids = layout_ids(prs)
    log.info("analyzing template", extra={"file": path.name, "slides": len(prs.slides)})

    colors = ColorObs()
    texts: list[TextObs] = []
    placed: list[Placed] = []
    titles: list[str] = []
    archetypes = []
    for n, slide in enumerate(prs.slides, start=1):
        texts.extend(inventory_slide(slide, resolver, colors))
        for shape in slide.shapes:
            box = shape_box(shape)
            if box is None:
                continue
            is_title = shape.is_placeholder and shape.placeholder_format.type in TITLE_TYPES
            zone = (
                Zone.TITLE
                if is_title
                else Zone.FOOTER
                if is_footer(shape, box, height)
                else Zone.CONTENT
            )
            placed.append(Placed(box, zone))
        arch = build_archetype(slide, n, ids[id(slide.slide_layout.element)], width, height)
        archetypes.append(arch)
        if arch.title:
            titles.append(arch.title)

    warnings: list[str] = []
    if len(masters) > 1:
        warnings.append(f"{len(masters)} slide masters found; theme taken from the first")
    if all(theme.colors.get(k) == v for k, v in OFFICE_2007_THEME.items()):
        warnings.append(
            "Theme uses the stock Office palette; the deck's real colors are direct "
            "formatting. Build from style.json `palette`, not theme slots."
        )
    pic_slides = [a.slide_number for a in archetypes if any("picture" in w for w in a.warnings)]
    if pic_slides:
        warnings.append(
            f"Slides {pic_slides} carry tables/charts as pictures; no native style to learn "
            "from them"
        )

    style = StyleSpec(
        source=source,
        slide_size=_aspect(width, height),
        theme_name=theme.name,
        theme_colors=theme.colors,
        theme_fonts=theme.fonts,
        master_text_styles=MasterTextStyles(
            title=resolver.master_levels("titleStyle", levels=1),
            body=resolver.master_levels("bodyStyle"),
        ),
        fonts=aggregate_fonts(texts),
        allowed_fonts=allowed_values(texts, "font"),
        allowed_font_sizes_pt=allowed_values(texts, "size_pt"),
        palette=Palette(
            text=color_usages(colors.text, theme.colors),
            fill=color_usages(colors.fill, theme.colors),
            line=color_usages(colors.line, theme.colors),
        ),
        grid=infer_grid(placed, width, height),
        text_rules=infer_text_rules(titles),
    )
    library = LayoutLibrary(
        source=source,
        layouts=build_layouts(prs, ids),
        slides=archetypes,
        warnings=warnings,
    )
    return style, library


def write_outputs(style: StyleSpec, library: LayoutLibrary, out_dir: Path) -> tuple[Path, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    style_path, layouts_path = out_dir / "style.json", out_dir / "layouts.json"
    style_path.write_text(style.model_dump_json(indent=2) + "\n", encoding="utf-8")
    layouts_path.write_text(library.model_dump_json(indent=2) + "\n", encoding="utf-8")
    return style_path, layouts_path
