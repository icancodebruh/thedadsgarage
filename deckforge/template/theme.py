"""Theme parsing and effective text-style resolution (run -> shape -> layout -> master)."""

from __future__ import annotations

from dataclasses import dataclass

from lxml import etree
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.presentation import Presentation
from pptx.shapes.base import BaseShape
from pptx.slide import SlideLayout, SlideMaster

from deckforge.template.models import TextLevelStyle, ThemeFonts
from deckforge.template.ooxml import DEFAULT_FONT_SIZE_PT, ColorResolver, Element, first, xpath

THEME_SLOTS = (
    "dk1", "lt1", "dk2", "lt2",
    "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
    "hlink", "folHlink",
)  # fmt: skip

TITLE_TYPES = frozenset({PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE})
OTHER_TYPES = frozenset({PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER})


@dataclass(frozen=True)
class Theme:
    name: str
    colors: dict[str, str]
    fonts: ThemeFonts
    text_defaults: Element | None  # <a:objectDefaults>/<a:txDef>/<a:lstStyle>


def parse_theme(master: SlideMaster) -> Theme:
    part = master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    colors: dict[str, str] = {}
    scheme = first(root, "a:themeElements/a:clrScheme")
    for slot in THEME_SLOTS:
        el = first(scheme, f"a:{slot}/*")
        if el is not None:
            value = el.get("val") if etree.QName(el).localname == "srgbClr" else el.get("lastClr")
            if value:
                colors[slot] = value.upper()

    def font(kind: str) -> str:
        el = first(root, f"a:themeElements/a:fontScheme/a:{kind}/a:latin")
        return el.get("typeface", "") if el is not None else ""

    return Theme(
        name=root.get("name", ""),
        colors=colors,
        fonts=ThemeFonts(major=font("majorFont"), minor=font("minorFont")),
        text_defaults=first(root, "a:objectDefaults/a:txDef/a:lstStyle"),
    )


def master_clr_map(master: SlideMaster) -> dict[str, str]:
    el = first(master.element, "p:clrMap")
    return dict(el.attrib) if el is not None else {}


@dataclass(frozen=True)
class RunStyle:
    font: str
    size_pt: float
    bold: bool
    color: str | None


class TextStyleResolver:
    """Resolves effective run properties by walking the OOXML inheritance chain."""

    def __init__(self, prs: Presentation, master: SlideMaster, theme: Theme) -> None:
        self.theme = theme
        self.colors = ColorResolver(theme.colors, master_clr_map(master))
        self._master = master
        self._tx_styles = first(master.element, "p:txStyles")
        self._pres_default = first(prs.element, "p:defaultTextStyle")

    # -- chain construction ---------------------------------------------------------------

    def _master_style(self, name: str) -> Element | None:
        return first(self._tx_styles, f"p:{name}")

    def level_chain(self, shape: BaseShape | None, layout: SlideLayout | None) -> list[Element]:
        """List-style containers (`a:lstStyle`-like) from most to least specific."""
        chain: list[Element] = []
        if shape is not None:
            own = first(shape.element, "p:txBody/a:lstStyle")
            if own is not None:
                chain.append(own)
        ph_type = _ph_type(shape)
        if ph_type is None:
            if self.theme.text_defaults is not None:
                chain.append(self.theme.text_defaults)
            if self._pres_default is not None:
                chain.append(self._pres_default)
            return chain
        assert shape is not None
        idx = shape.placeholder_format.idx
        layout_ph = (
            find_matching_placeholder(layout.placeholders, idx, ph_type)
            if layout is not None
            else None
        )
        if layout_ph is not None:
            chain.extend(xpath(layout_ph.element, "p:txBody/a:lstStyle"))
        master_ph = find_matching_placeholder(self._master.placeholders, None, ph_type)
        if master_ph is not None:
            chain.extend(xpath(master_ph.element, "p:txBody/a:lstStyle"))
        style = (
            "titleStyle"
            if ph_type in TITLE_TYPES
            else "otherStyle"
            if ph_type in OTHER_TYPES
            else "bodyStyle"
        )
        master_style = self._master_style(style)
        if master_style is not None:
            chain.append(master_style)
        if self._pres_default is not None:
            chain.append(self._pres_default)
        return chain

    # -- property lookup ------------------------------------------------------------------

    def resolve(self, r_pr: Element | None, p_pr: Element | None, chain: list[Element]) -> RunStyle:
        level = int(p_pr.get("lvl", "0")) + 1 if p_pr is not None else 1
        candidates: list[Element] = [r_pr] if r_pr is not None else []
        for container in chain:
            d = first(container, f"a:lvl{level}pPr/a:defRPr")
            if d is not None:
                candidates.append(d)
        return self._style_from(candidates)

    def _style_from(self, rprs: list[Element]) -> RunStyle:
        size = next((int(r.get("sz")) / 100 for r in rprs if r.get("sz")), DEFAULT_FONT_SIZE_PT)
        bold = next((r.get("b") in ("1", "true") for r in rprs if r.get("b") is not None), False)
        font = next(
            (lat.get("typeface") for r in rprs if (lat := first(r, "a:latin")) is not None),
            "+mn-lt",
        )
        color = next(
            (c for r in rprs if (c := self.colors.solid_fill(r)) is not None),
            None,
        )
        return RunStyle(font=self.font_name(font), size_pt=size, bold=bold, color=color)

    def font_name(self, typeface: str | None) -> str:
        if typeface in (None, "", "+mn-lt", "+mn-ea", "+mn-cs"):
            return self.theme.fonts.minor
        if typeface in ("+mj-lt", "+mj-ea", "+mj-cs"):
            return self.theme.fonts.major
        return typeface

    def master_levels(self, name: str, levels: int = 5) -> list[TextLevelStyle]:
        container = self._master_style(name)
        chain = [c for c in (container, self._pres_default) if c is not None]
        out: list[TextLevelStyle] = []
        for lvl in range(1, levels + 1):
            p_pr = first(container, f"a:lvl{lvl}pPr")
            rprs = [d for c in chain if (d := first(c, f"a:lvl{lvl}pPr/a:defRPr")) is not None]
            style = self._style_from(rprs)
            bullet = first(p_pr, "a:buChar")
            out.append(
                TextLevelStyle(
                    level=lvl,
                    font=style.font,
                    size_pt=style.size_pt,
                    color=style.color,
                    bold=style.bold,
                    align=p_pr.get("algn") if p_pr is not None else None,
                    bullet=bullet.get("char") if bullet is not None else None,
                )
            )
        return out


def _ph_type(shape: BaseShape | None) -> PP_PLACEHOLDER | None:
    if shape is None or not shape.is_placeholder:
        return None
    return shape.placeholder_format.type


def find_matching_placeholder(
    placeholders: object, idx: int | None, ph_type: PP_PLACEHOLDER
) -> BaseShape | None:
    """Match a placeholder by idx first, then by type (title types are interchangeable)."""
    candidates = list(placeholders)  # type: ignore[call-overload]
    if idx is not None:
        for ph in candidates:
            if ph.placeholder_format.idx == idx:
                return ph  # type: ignore[no-any-return]
    wanted = TITLE_TYPES if ph_type in TITLE_TYPES else {ph_type}
    for ph in candidates:
        if ph.placeholder_format.type in wanted:
            return ph  # type: ignore[no-any-return]
    # Body-like placeholders (obj, subTitle) inherit from the master body placeholder.
    if ph_type not in TITLE_TYPES | OTHER_TYPES:
        for ph in candidates:
            if ph.placeholder_format.type == PP_PLACEHOLDER.BODY:
                return ph  # type: ignore[no-any-return]
    return None
